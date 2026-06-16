"""Reusable OpenPyXL styles for the MIS workbook."""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INR_NUMBER_FORMAT = '₹#,##0;[Red](₹#,##0);-'
PERCENT_NUMBER_FORMAT = '0.0%;[Red](0.0%);-'
TEXT_NUMBER_FORMAT = "@"

DARK_BLUE = "FF1F2D4E"
MID_BLUE = "FF2F5496"
SECTION_BLUE = "FF4472C4"
TOTAL_BLUE = "FFD6E4F7"
ALT_ROW = "FFEBF3FB"
WHITE = "FFFFFFFF"
SM_FILL = "FFE3F2FD"
RD_FILL = "FFE8F5E9"
GA_FILL = "FFF3E5F5"
COGS_FILL = "FFFCE4EC"
YELLOW_FILL = "FFFFF9C4"
LIGHT_GRAY = "FFF5F5F5"

THIN_GRAY = Side(style="thin", color="FFD0D7E2")
MEDIUM_BLUE = Side(style="medium", color=DARK_BLUE)


def base_alignment(horizontal: str = "right", wrap_text: bool = False) -> Alignment:
    """Return a standard workbook alignment."""
    return Alignment(horizontal=horizontal, vertical="center", wrap_text=wrap_text)


STYLE_LIBRARY = {
    "title": {
        "fill": PatternFill("solid", fgColor=DARK_BLUE),
        "font": Font(color="FFFFFFFF", bold=True, size=13),
        "alignment": base_alignment(horizontal="left"),
    },
    "subtitle": {
        "fill": PatternFill("solid", fgColor=MID_BLUE),
        "font": Font(color="FFFFFFFF", italic=True, size=10),
        "alignment": base_alignment(horizontal="left"),
    },
    "table_header": {
        "fill": PatternFill("solid", fgColor=DARK_BLUE),
        "font": Font(color="FFFFFFFF", bold=True),
        "alignment": base_alignment(horizontal="center", wrap_text=True),
    },
    "section": {
        "fill": PatternFill("solid", fgColor=SECTION_BLUE),
        "font": Font(color="FFFFFFFF", bold=True),
        "alignment": base_alignment(horizontal="left"),
    },
    "total": {
        "fill": PatternFill("solid", fgColor=TOTAL_BLUE),
        "font": Font(color="FF000000", bold=True),
        "alignment": base_alignment(horizontal="right"),
    },
    "dark_total": {
        "fill": PatternFill("solid", fgColor=DARK_BLUE),
        "font": Font(color="FFFFFFFF", bold=True),
        "alignment": base_alignment(horizontal="right"),
    },
    "alt_row": {
        "fill": PatternFill("solid", fgColor=ALT_ROW),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "plain_row": {
        "fill": PatternFill("solid", fgColor=WHITE),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "sm_row": {
        "fill": PatternFill("solid", fgColor=SM_FILL),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "rd_row": {
        "fill": PatternFill("solid", fgColor=RD_FILL),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "ga_row": {
        "fill": PatternFill("solid", fgColor=GA_FILL),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "cogs_row": {
        "fill": PatternFill("solid", fgColor=COGS_FILL),
        "font": Font(color="FF000000"),
        "alignment": base_alignment(horizontal="right"),
    },
    "metric": {
        "fill": PatternFill("solid", fgColor=ALT_ROW),
        "font": Font(color="FF000000", bold=True),
        "alignment": base_alignment(horizontal="left"),
    },
    "metric_highlight": {
        "fill": PatternFill("solid", fgColor=YELLOW_FILL),
        "font": Font(color="FF000000", bold=True),
        "alignment": base_alignment(horizontal="left"),
    },
    "note": {
        "font": Font(color="FF000000", italic=True),
        "alignment": base_alignment(horizontal="left", wrap_text=True),
    },
    "gray_total": {
        "fill": PatternFill("solid", fgColor=LIGHT_GRAY),
        "font": Font(color="FF000000", bold=True),
        "alignment": base_alignment(horizontal="right"),
    },
}


def apply_style(cell, style_name: str, number_format: str | None = None) -> None:
    """Apply a named style to one cell."""
    style = STYLE_LIBRARY[style_name]
    if "fill" in style:
        cell.fill = copy(style["fill"])
    if "font" in style:
        cell.font = copy(style["font"])
    if "alignment" in style:
        cell.alignment = copy(style["alignment"])
    cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
    if number_format:
        cell.number_format = number_format


def apply_row_style(
    worksheet,
    row_number: int,
    start_column: int,
    end_column: int,
    style_name: str,
    number_format: str | None = None,
) -> None:
    """Apply the same style across a row range."""
    for column_number in range(start_column, end_column + 1):
        apply_style(worksheet.cell(row_number, column_number), style_name, number_format=number_format)


def apply_label_cell(cell, style_name: str, indent: int = 0) -> None:
    """Style a row label cell with left alignment."""
    apply_style(cell, style_name)
    cell.alignment = base_alignment(horizontal="left", wrap_text=True)
    if indent:
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
            indent=indent,
        )


def apply_title_band(worksheet, row_number: int, start_column: int, end_column: int, title: str, style_name: str) -> None:
    """Merge and style a title row."""
    worksheet.merge_cells(start_row=row_number, start_column=start_column, end_row=row_number, end_column=end_column)
    cell = worksheet.cell(row_number, start_column)
    cell.value = title
    apply_style(cell, style_name)
    cell.border = Border(top=MEDIUM_BLUE, bottom=MEDIUM_BLUE, left=MEDIUM_BLUE, right=MEDIUM_BLUE)


def apply_border_band(worksheet, row_number: int, start_column: int, end_column: int) -> None:
    """Add a medium top and bottom border for important totals."""
    for column_number in range(start_column, end_column + 1):
        cell = worksheet.cell(row_number, column_number)
        cell.border = Border(top=MEDIUM_BLUE, bottom=MEDIUM_BLUE, left=THIN_GRAY, right=THIN_GRAY)
