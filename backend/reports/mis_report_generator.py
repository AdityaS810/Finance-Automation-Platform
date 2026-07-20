"""Generate the company MIS workbook from BigQuery and Zoho detail data."""

from __future__ import annotations

import json
import os
from copy import copy
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.reports.report_periods import (
    get_consolidated_report_period,
    get_fy_label,
    get_month_periods,
    get_selected_report_period,
    parse_financial_year_start,
)


DEFAULT_PROJECT_ID = "internal-project-work-497507"
MIS_MONTHLY_PL_VIEW = "finance_gold.mis_monthly_pl"
DASHBOARD_SUMMARY_VIEW = "finance_gold.dashboard_summary"
BRONZE_RAW_VIEW = "finance_bronze.zoho_raw"
INVOICES_VIEW = "finance_silver.fact_invoices"
BILLS_VIEW = "finance_silver.fact_bills"
FX_RATE_DEFAULT = 90.0

MONTHLY_TEMPLATE_SHEET = "Monthly P&L"
QUARTERLY_TEMPLATE_SHEET = "Quarterly P&L"
COGS_TEMPLATE_SHEET = "COGS Allocation Working"
TEMPLATE_FILE_NAME = "MidofficeData_KeyMetrics_PL_FY2526.xlsx"
TEMPLATE_FILE_PATTERN = "MidofficeData_KeyMetrics_PL_*.xlsx"

MONTHLY_ROW_LABELS = {
    4: "REVENUE",
    5: "Client Revenue - Professional Services",
    6: "Client Revenue - Product",
    7: "Other Income",
    8: "TOTAL REVENUE",
    9: "COGS",
    10: "Offshore COGS",
    11: "Onshore COGS",
    12: "Technology Costs",
    13: "TOTAL COGS",
    14: "GROSS PROFIT",
    15: "S&M",
    16: "Onshore Consultant Experience",
    17: "Advertising & Marketing",
    18: "Travel Expenses",
    19: "Meals & Entertainment",
    20: "S&M Total",
    21: "R&D",
    22: "R&D Salaries",
    23: "Consultant & Contractor Expense",
    24: "Software Subs",
    25: "R&D Total",
    26: "G&A",
    27: "Office Rent",
    28: "IT & Internet",
    29: "Legal",
    30: "Audit & Non-Op",
    31: "Other G&A",
    32: "G&A Total",
    33: "TOTAL OPEX",
    34: "NET PROFIT/(LOSS)",
}

QUARTERLY_ROW_LABELS = {
    6: "  Client Revenue - Professional Services",
    7: "  Client Revenue - Product",
    8: "  Other Income",
    12: "  Offshore COGS",
    13: "  Onshore COGS",
    14: "",
    15: "",
    16: "",
    25: "      Onshore Consultant Experience",
    33: "      Consultant & Contractor Expense",
    38: "      Office Rent",
}

MONTHLY_PREVIEW_ITEMS = [
    ("Client Revenue - Professional Services", "techm_billings"),
    ("Client Revenue - Product", "bsm_revenue"),
    ("Other Income", "fd_interest"),
    ("Offshore COGS", "offshore_cogs"),
    ("Onshore COGS", "onshore_cogs"),
    ("Technology Costs", "technology_costs"),
    ("COGS Total", "cogs_total"),
    ("Onshore Consultant Experience", "ramki_sm"),
    ("Advertising & Marketing", "advertising_marketing"),
    ("Travel Expenses", "travel_expenses"),
    ("Meals & Entertainment", "meals_entertainment"),
    ("R&D Salaries", "rd_salaries"),
    ("Consultant & Contractor Expense", "consultant_expense"),
    ("Software Subs", "software_subscriptions"),
    ("Office Rent", "rent"),
    ("IT & Internet", "it_internet"),
    ("Legal", "legal"),
    ("Audit & Non-Op", "audit_non_operating"),
    ("Other G&A", "other_ga"),
]

EXCEL_ROW_LABEL_TO_RULE_KEY = {label.strip().lower(): rule_key for label, rule_key in MONTHLY_PREVIEW_ITEMS}
EXCEL_ROW_LABEL_TO_RULE_KEY["other g&a"] = "delivery_india"


load_dotenv()


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _template_path() -> Path:
    templates_dir = _repo_root() / "templates"
    exact_match = templates_dir / TEMPLATE_FILE_NAME
    if exact_match.exists():
        return exact_match

    template_matches = sorted(templates_dir.glob(TEMPLATE_FILE_PATTERN))
    return template_matches[0] if template_matches else exact_match


def _mis_mapping_path() -> Path:
    return _repo_root() / "config" / "mis_mapping.yaml"


def _cogs_config_path() -> Path:
    return _repo_root() / "config" / "cogs_allocation.yaml"


def _get_project_id(project_id: str | None = None) -> str:
    """Use an explicit project id, then environment, then the known project."""
    return project_id or os.getenv("GCP_PROJECT_ID") or DEFAULT_PROJECT_ID


def _table_name(project_id: str, view_name: str) -> str:
    """Build a fully qualified BigQuery table or view name."""
    return f"`{project_id}.{view_name}`"


def _normalise_org_filter(org_filter: str | None = None) -> str:
    """Map UI labels and missing values to safe organization keys."""
    if not org_filter:
        return "all"
    normalised = str(org_filter).strip().lower()
    aliases = {
        "all organizations": "all",
        "all": "all",
        "india": "india",
        "in": "india",
        "us": "us",
        "u.s.": "us",
        "usa": "us",
        "u.s.a": "us",
    }
    return aliases.get(normalised, "all")


def _financial_year_token(financial_year_start: int | str) -> str:
    """Return a filename token like FY2025_26."""
    start_year = parse_financial_year_start(financial_year_start)
    return f"FY{start_year}_{str(start_year + 1)[-2:]}"


def _utc_now_naive() -> datetime:
    """Return a timezone-naive UTC timestamp that Excel can store."""
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _format_currency(value: float | int | None) -> str:
    """Return a compact INR value for Streamlit metric cards."""
    amount = float(value or 0)
    if abs(amount) >= 10_000_000:
        return f"INR {amount / 10_000_000:.2f} Cr"
    if abs(amount) >= 100_000:
        return f"INR {amount / 100_000:.2f} L"
    return f"INR {amount:,.0f}"


def _format_inr_short(value: float | int | None) -> str:
    """Return a human-readable rupee string for note blocks."""
    amount = float(value or 0)
    if abs(amount) >= 10_000_000:
        return f"₹{amount / 10_000_000:.2f} Cr"
    if abs(amount) >= 100_000:
        return f"₹{amount / 100_000:.1f}L"
    return f"₹{amount:,.0f}"


def _first_value(dataframe: pd.DataFrame, column_name: str, default: Any = 0) -> Any:
    """Read one value from a one-row dataframe without raising on empty data."""
    if dataframe.empty or column_name not in dataframe.columns:
        return default

    value = dataframe.iloc[0][column_name]
    if pd.isna(value):
        return default

    return value


def _safe_sum(dataframe: pd.DataFrame, column_name: str) -> float:
    """Sum a numeric dataframe column and treat missing values as zero."""
    if dataframe.empty or column_name not in dataframe.columns:
        return 0.0
    return float(pd.to_numeric(dataframe[column_name], errors="coerce").fillna(0).sum())


def _numeric_column(dataframe: pd.DataFrame, column_name: str) -> pd.Series:
    """Return a numeric series, or a zero-filled series when a column is absent."""
    if column_name not in dataframe.columns:
        return pd.Series(0, index=dataframe.index, dtype="float64")
    return pd.to_numeric(dataframe[column_name], errors="coerce").fillna(0)


def _text_column(dataframe: pd.DataFrame, column_name: str, default: str = "") -> pd.Series:
    """Return a text series, or a default-filled series when a column is absent."""
    if column_name not in dataframe.columns:
        return pd.Series(default, index=dataframe.index, dtype="object")
    return dataframe[column_name].fillna(default).astype(str)


def _safe_ratio(numerator: float, denominator: float) -> float:
    """Return a ratio, or zero when the denominator is not usable."""
    if not denominator:
        return 0.0
    return numerator / denominator


def _query_to_dataframe(client, query: str) -> pd.DataFrame:
    """Run a BigQuery query and convert rows to pandas without extra packages."""
    result = client.query(query).result()
    columns = [field.name for field in result.schema]
    rows = [dict(row.items()) for row in result]
    return pd.DataFrame(rows, columns=columns)


def _read_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        return yaml.safe_load(stream) or {}


def _is_active_mapping(value: Any) -> bool:
    return str(value).strip().lower() in {"active", "true", "yes", "1"}


def _org_mapping_applies(mapping_org: Any, report_org_filter: str | None) -> bool:
    mapping_value = _normalise_org_filter(_to_string(mapping_org))
    report_value = _normalise_org_filter(report_org_filter)
    return mapping_value == "all" or report_value == "all" or mapping_value == report_value


def _rule_key_for_excel_label(label: Any) -> str:
    cleaned_label = _to_string(label).strip()
    return EXCEL_ROW_LABEL_TO_RULE_KEY.get(cleaned_label.lower()) or cleaned_label.lower().replace("&", "and").replace("/", " ").replace(" ", "_")


def _merge_account_mapping_rules(
    existing_rules: dict[str, Any],
    account_mappings: list[dict[str, Any]] | None,
    section_name: str,
    org_filter: str | None,
) -> dict[str, Any]:
    """Merge active account-code mappings into legacy keyword rules."""
    merged_rules = {rule_key: dict(rule or {}) for rule_key, rule in existing_rules.items()}
    for mapping in account_mappings or []:
        if not isinstance(mapping, dict) or not mapping.get("active"):
            continue
        if _to_string(mapping.get("excel_section")).strip().lower() != section_name:
            continue
        if not _org_mapping_applies(mapping.get("organization"), org_filter):
            continue

        account_id = _to_string(mapping.get("zoho_account_id")).strip()
        account_code = _to_string(mapping.get("zoho_account_code")).strip()
        if not account_id and not account_code:
            continue

        rule_key = _rule_key_for_excel_label(mapping.get("excel_row_label"))
        rule = dict(merged_rules.get(rule_key, {}))
        source_types = set(rule.get("source_types") or ["invoice", "bill", "journal"])
        match_fields = set(rule.get("match_fields") or [])
        match_fields.update(["account_id", "account_code", "account_name"])
        account_ids = set(rule.get("account_ids") or [])
        account_codes = set(rule.get("account_codes") or [])
        if account_id:
            account_ids.add(account_id)
        if account_code:
            account_codes.add(account_code)

        rule.update(
            {
                "source_types": sorted(source_types),
                "match_fields": sorted(match_fields),
                "account_ids": sorted(account_ids),
                "account_codes": sorted(account_codes),
            }
        )
        merged_rules[rule_key] = rule
    return merged_rules


def _parse_financial_year(financial_year: str) -> tuple[int, int]:
    start_year = parse_financial_year_start(financial_year)
    return start_year, start_year + 1


def _financial_year_months(financial_year: str, configured_months: list[str] | None = None) -> list[str]:
    del configured_months
    start_year, _ = _parse_financial_year(financial_year)
    return [period["key"] for period in get_month_periods(start_year)]


def _month_labels(months: list[str]) -> dict[str, str]:
    labels = {}
    for month_key in months:
        month_date = pd.Timestamp(f"{month_key}-01")
        labels[month_key] = month_date.strftime("%b %y")
    return labels


def _quarter_months(months: list[str]) -> dict[str, list[str]]:
    return {
        "Q1": months[0:3],
        "Q2": months[3:6],
        "Q3": months[6:9],
        "Q4": months[9:12],
    }


def _month_column_map(start_column_index: int, months: list[str]) -> dict[str, int]:
    return {month_key: start_column_index + index for index, month_key in enumerate(months)}


def _zero_month_map(months: list[str]) -> dict[str, float]:
    return {month_key: 0.0 for month_key in months}


def _round_currency(value: float) -> float:
    return round(float(value or 0), 2)


def _normalize_month(value: Any) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.strftime("%Y-%m")


def _format_month_key(month_key: str | None) -> str:
    """Convert YYYY-MM to a readable month label."""
    if not month_key:
        return "-"
    parsed = pd.to_datetime(f"{month_key}-01", errors="coerce")
    if pd.isna(parsed):
        return str(month_key)
    return parsed.strftime("%b %Y")


def _date_to_sql(value: date | datetime) -> str:
    """Render a date literal for BigQuery."""
    return (value.date() if isinstance(value, datetime) else value).strftime("%Y-%m-%d")


def _filter_dataframe_by_date_range(
    dataframe: pd.DataFrame,
    date_column: str,
    start_date: date,
    end_date: date,
) -> pd.DataFrame:
    """Keep only rows inside the requested date range."""
    if dataframe.empty or date_column not in dataframe.columns:
        return dataframe.copy()

    parsed_dates = pd.to_datetime(dataframe[date_column], errors="coerce")
    filtered = dataframe.loc[
        parsed_dates.notna()
        & (parsed_dates.dt.date >= start_date)
        & (parsed_dates.dt.date <= end_date)
    ].copy()
    return filtered.reset_index(drop=True)


def _to_string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _parse_raw_payload(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return {}
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}
    return {}


def _get_line_items(raw_payload: dict[str, Any]) -> list[dict[str, Any]]:
    line_items = raw_payload.get("line_items")
    if isinstance(line_items, list):
        return [item for item in line_items if isinstance(item, dict)]
    return []


def _extract_amount_from_line_item(line_item: dict[str, Any]) -> float:
    for key in ("item_total", "total", "line_total", "amount", "net_amount", "total_amount"):
        amount = _to_float(line_item.get(key), default=0.0)
        if amount:
            return amount

    rate = _to_float(line_item.get("rate"), default=0.0)
    quantity = _to_float(line_item.get("quantity"), default=1.0)
    if rate:
        return rate * quantity

    return 0.0


def _convert_to_inr(amount: float, currency_code: str | None, exchange_rate: Any, fx_rate_default: float) -> float:
    normalized_amount = _to_float(amount, default=0.0)
    if not normalized_amount:
        return 0.0

    currency = _to_string(currency_code).strip().upper()
    rate = _to_float(exchange_rate, default=0.0)
    if currency in {"", "INR"}:
        return normalized_amount
    if currency in {"USD", "US$"}:
        return normalized_amount * (rate or fx_rate_default)
    if rate and not currency:
        return normalized_amount * rate
    return normalized_amount


def fetch_gold_mis_data(
    project_id: str | None = None,
    org_filter: str | None = None,
    report_period: dict[str, Any] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Query the Consume MIS monthly view for the requested reporting window."""
    from google.cloud import bigquery

    resolved_project_id = _get_project_id(project_id)
    selected_org_key = _normalise_org_filter(org_filter)
    client = bigquery.Client(project=resolved_project_id)
    filters: list[str] = []
    if selected_org_key != "all":
        filters.append(f"source_org_key = '{selected_org_key}'")
    if report_period:
        filters.append(
            "DATE(report_month) BETWEEN "
            f"DATE '{_date_to_sql(report_period['start_date'])}' AND DATE '{_date_to_sql(report_period['end_date'])}'"
        )
    where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""

    monthly_query = f"""
        SELECT *
        FROM {_table_name(resolved_project_id, MIS_MONTHLY_PL_VIEW)}
        {where_clause}
        ORDER BY report_month
    """

    monthly_pl_df = _query_to_dataframe(client, monthly_query)
    return monthly_pl_df, pd.DataFrame()


def fetch_detailed_mis_data(
    project_id: str | None = None,
    org_filter: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Fetch invoice, bill, and journal raw payloads for the selected reporting period."""
    from google.cloud import bigquery

    resolved_project_id = _get_project_id(project_id)
    selected_org_key = _normalise_org_filter(org_filter)
    client = bigquery.Client(project=resolved_project_id)
    org_filter_sql = f"AND source_org_key = '{selected_org_key}'" if selected_org_key != "all" else ""
    date_filter_sql = ""
    if start_date and end_date:
        start_sql = _date_to_sql(start_date)
        end_sql = _date_to_sql(end_date)
        date_filter_sql = f"AND SAFE_CAST(JSON_VALUE(raw_json, '$.date') AS DATE) BETWEEN DATE '{start_sql}' AND DATE '{end_sql}'"

    invoice_query = f"""
        WITH latest AS (
            SELECT
                run_id,
                source_record_id,
                loaded_at,
                raw_json
            FROM {_table_name(resolved_project_id, BRONZE_RAW_VIEW)}
            WHERE entity_name = 'invoices'
              {org_filter_sql}
              {date_filter_sql}
            QUALIFY ROW_NUMBER() OVER (PARTITION BY source_record_id ORDER BY loaded_at DESC) = 1
        )
        SELECT
            run_id,
            source_record_id AS invoice_id,
            JSON_VALUE(raw_json, '$.invoice_number') AS invoice_number,
            JSON_VALUE(raw_json, '$.customer_name') AS customer_name,
            SAFE_CAST(JSON_VALUE(raw_json, '$.date') AS DATE) AS invoice_date,
            JSON_VALUE(raw_json, '$.currency_code') AS currency_code,
            SAFE_CAST(JSON_VALUE(raw_json, '$.exchange_rate') AS NUMERIC) AS exchange_rate,
            SAFE_CAST(JSON_VALUE(raw_json, '$.total') AS NUMERIC) AS total_amount,
            JSON_VALUE(raw_json, '$.reference_number') AS reference_number,
            JSON_VALUE(raw_json, '$.notes') AS notes,
            raw_json,
            loaded_at
        FROM latest
    """
    bill_query = f"""
        WITH latest AS (
            SELECT
                run_id,
                source_record_id,
                loaded_at,
                raw_json
            FROM {_table_name(resolved_project_id, BRONZE_RAW_VIEW)}
            WHERE entity_name = 'bills'
              {org_filter_sql}
              {date_filter_sql}
            QUALIFY ROW_NUMBER() OVER (PARTITION BY source_record_id ORDER BY loaded_at DESC) = 1
        )
        SELECT
            run_id,
            source_record_id AS bill_id,
            JSON_VALUE(raw_json, '$.bill_number') AS bill_number,
            JSON_VALUE(raw_json, '$.vendor_name') AS vendor_name,
            SAFE_CAST(JSON_VALUE(raw_json, '$.date') AS DATE) AS bill_date,
            JSON_VALUE(raw_json, '$.currency_code') AS currency_code,
            SAFE_CAST(JSON_VALUE(raw_json, '$.exchange_rate') AS NUMERIC) AS exchange_rate,
            SAFE_CAST(JSON_VALUE(raw_json, '$.total') AS NUMERIC) AS total_amount,
            JSON_VALUE(raw_json, '$.notes') AS notes,
            raw_json,
            loaded_at
        FROM latest
    """
    journal_query = f"""
        WITH latest AS (
            SELECT
                run_id,
                source_record_id,
                loaded_at,
                raw_json
            FROM {_table_name(resolved_project_id, BRONZE_RAW_VIEW)}
            WHERE entity_name = 'journals'
              {org_filter_sql}
              {date_filter_sql}
            QUALIFY ROW_NUMBER() OVER (PARTITION BY source_record_id ORDER BY loaded_at DESC) = 1
        )
        SELECT
            run_id,
            source_record_id AS journal_id,
            JSON_VALUE(raw_json, '$.journal_number') AS journal_number,
            SAFE_CAST(JSON_VALUE(raw_json, '$.date') AS DATE) AS journal_date,
            JSON_VALUE(raw_json, '$.reference_number') AS reference_number,
            JSON_VALUE(raw_json, '$.notes') AS notes,
            JSON_VALUE(raw_json, '$.currency_code') AS currency_code,
            SAFE_CAST(JSON_VALUE(raw_json, '$.exchange_rate') AS NUMERIC) AS exchange_rate,
            COALESCE(
                SAFE_CAST(JSON_VALUE(raw_json, '$.total') AS NUMERIC),
                SAFE_CAST(JSON_VALUE(raw_json, '$.amount') AS NUMERIC)
            ) AS total_amount,
            raw_json,
            loaded_at
        FROM latest
    """

    return (
        _query_to_dataframe(client, invoice_query),
        _query_to_dataframe(client, bill_query),
        _query_to_dataframe(client, journal_query),
    )


def _candidate_text_map(row: pd.Series, raw_payload: dict[str, Any], line_item: dict[str, Any] | None = None) -> dict[str, str]:
    line_item = line_item or {}
    return {
        "account_id": " ".join(
            filter(
                None,
                [
                    _to_string(line_item.get("account_id")),
                    _to_string(line_item.get("accountId")),
                    _to_string(raw_payload.get("account_id")),
                    _to_string(row.get("account_id")),
                ],
            )
        ),
        "account_code": " ".join(
            filter(
                None,
                [
                    _to_string(line_item.get("account_code")),
                    _to_string(line_item.get("accountCode")),
                    _to_string(raw_payload.get("account_code")),
                    _to_string(row.get("account_code")),
                ],
            )
        ),
        "counterparty_name": " ".join(
            filter(
                None,
                [
                    _to_string(row.get("customer_name")),
                    _to_string(row.get("vendor_name")),
                    _to_string(raw_payload.get("customer_name")),
                    _to_string(raw_payload.get("vendor_name")),
                    _to_string(raw_payload.get("contact_name")),
                ],
            )
        ),
        "description": " ".join(
            filter(
                None,
                [
                    _to_string(line_item.get("description")),
                    _to_string(raw_payload.get("description")),
                ],
            )
        ),
        "item_name": " ".join(
            filter(
                None,
                [
                    _to_string(line_item.get("item_name")),
                    _to_string(line_item.get("name")),
                    _to_string(line_item.get("product_name")),
                    _to_string(line_item.get("line_item_name")),
                ],
            )
        ),
        "notes": " ".join(
            filter(
                None,
                [
                    _to_string(row.get("notes")),
                    _to_string(raw_payload.get("notes")),
                    _to_string(raw_payload.get("memo")),
                ],
            )
        ),
        "reference_number": " ".join(
            filter(
                None,
                [
                    _to_string(row.get("reference_number")),
                    _to_string(row.get("invoice_number")),
                    _to_string(row.get("bill_number")),
                    _to_string(row.get("journal_number")),
                    _to_string(raw_payload.get("reference_number")),
                ],
            )
        ),
        "account_name": " ".join(
            filter(
                None,
                [
                    _to_string(line_item.get("account_name")),
                    _to_string(line_item.get("account")),
                    _to_string(line_item.get("name")),
                ],
            )
        ),
    }


def _match_rule(text_map: dict[str, str], rule: dict[str, Any]) -> bool:
    account_ids = {str(account_id).strip().lower() for account_id in rule.get("account_ids", []) if str(account_id).strip()}
    if account_ids:
        candidate_ids = {part.strip().lower() for part in text_map.get("account_id", "").split() if part.strip()}
        if account_ids.intersection(candidate_ids):
            return True

    account_codes = {str(account_code).strip().lower() for account_code in rule.get("account_codes", []) if str(account_code).strip()}
    if account_codes:
        candidate_codes = {part.strip().lower() for part in text_map.get("account_code", "").split() if part.strip()}
        if account_codes.intersection(candidate_codes):
            return True

    match_fields = rule.get("match_fields", [])
    haystack = " ".join(text_map.get(field, "") for field in match_fields).lower()
    if not haystack:
        return False

    exclude_keywords = [keyword.lower() for keyword in rule.get("exclude_keywords", [])]
    if any(keyword in haystack for keyword in exclude_keywords):
        return False

    keywords = [keyword.lower() for keyword in rule.get("keywords", [])]
    return any(keyword in haystack for keyword in keywords)


def _classify_candidate(text_map: dict[str, str], source_type: str, rules: dict[str, Any]) -> str | None:
    for rule_key, rule in rules.items():
        if source_type not in rule.get("source_types", []):
            continue
        if _match_rule(text_map, rule):
            return rule_key
    return None


def _iter_detail_candidates(
    dataframe: pd.DataFrame,
    source_type: str,
    date_column: str,
    amount_column: str,
    valid_months: set[str],
    fx_rate_default: float,
):
    if dataframe.empty:
        return

    for _, row in dataframe.iterrows():
        raw_payload = _parse_raw_payload(row.get("raw_json"))
        month_key = _normalize_month(row.get(date_column) or raw_payload.get("date"))
        if month_key not in valid_months:
            continue

        currency_code = _to_string(row.get("currency_code") or raw_payload.get("currency_code"))
        exchange_rate = row.get("exchange_rate") if row.get("exchange_rate") is not None else raw_payload.get("exchange_rate")
        header_amount = _convert_to_inr(row.get(amount_column), currency_code, exchange_rate, fx_rate_default)

        line_items = _get_line_items(raw_payload)
        if source_type in {"bill", "journal"} and line_items:
            yielded_line = False
            for line_item in line_items:
                line_amount = _extract_amount_from_line_item(line_item)
                if not line_amount and len(line_items) == 1:
                    line_amount = _to_float(row.get(amount_column), default=0.0)
                if not line_amount:
                    continue

                yielded_line = True
                yield {
                    "month": month_key,
                    "source_type": source_type,
                    "amount_inr": _convert_to_inr(line_amount, currency_code, exchange_rate, fx_rate_default),
                    "text_map": _candidate_text_map(row, raw_payload, line_item),
                }

            if yielded_line:
                continue

        yield {
            "month": month_key,
            "source_type": source_type,
            "amount_inr": header_amount,
            "text_map": _candidate_text_map(row, raw_payload, None),
        }


def _aggregate_rule_totals(
    months: list[str],
    fx_rate_default: float,
    invoices_df: pd.DataFrame,
    bills_df: pd.DataFrame,
    journals_df: pd.DataFrame,
    rules: dict[str, Any],
) -> dict[str, dict[str, float]]:
    totals = {rule_key: _zero_month_map(months) for rule_key in rules}
    valid_months = set(months)

    for candidate in _iter_detail_candidates(invoices_df, "invoice", "invoice_date", "total_amount", valid_months, fx_rate_default) or []:
        category = _classify_candidate(candidate["text_map"], candidate["source_type"], rules)
        if category:
            totals[category][candidate["month"]] += candidate["amount_inr"]

    for candidate in _iter_detail_candidates(bills_df, "bill", "bill_date", "total_amount", valid_months, fx_rate_default) or []:
        category = _classify_candidate(candidate["text_map"], candidate["source_type"], rules)
        if category:
            totals[category][candidate["month"]] += candidate["amount_inr"]

    for candidate in _iter_detail_candidates(journals_df, "journal", "journal_date", "total_amount", valid_months, fx_rate_default) or []:
        category = _classify_candidate(candidate["text_map"], candidate["source_type"], rules)
        if category:
            totals[category][candidate["month"]] += candidate["amount_inr"]

    for rule_months in totals.values():
        for month_key, value in list(rule_months.items()):
            rule_months[month_key] = _round_currency(value)

    return totals


def _aggregate_cogs_detail(
    months: list[str],
    fx_rate_default: float,
    bills_df: pd.DataFrame,
    journals_df: pd.DataFrame,
    company_rules: dict[str, Any],
) -> dict[str, dict[str, float]]:
    cogs_rows = {
        "parashar": _zero_month_map(months),
        "sangeeth": _zero_month_map(months),
        "jitin": _zero_month_map(months),
        "dipti": _zero_month_map(months),
        "dhanashri": _zero_month_map(months),
        "pinkesh": _zero_month_map(months),
        "ramki": _zero_month_map(months),
        "dorothea": _zero_month_map(months),
        "bsm_crew": _zero_month_map(months),
        "conam_tech": _zero_month_map(months),
        "external_vendors": _zero_month_map(months),
        "technology_costs": _zero_month_map(months),
    }

    row_rules = {
        "parashar": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["parashar"]},
        "sangeeth": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["sangeeth"]},
        "jitin": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["jitin"]},
        "dipti": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["dipti"]},
        "dhanashri": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["dhanashri"]},
        "pinkesh": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["pinkesh"]},
        "dorothea": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["dorothea"]},
        "bsm_crew": {"source_types": ["bill", "journal"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["kashish", "priyanka", "nitin chandani"]},
        "conam_tech": {"source_types": ["bill"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["conam"]},
        "external_vendors": {"source_types": ["bill"], "match_fields": ["counterparty_name", "description", "item_name", "notes", "account_name"], "keywords": ["biltzentech", "altysys", "external delivery"]},
    }

    valid_months = set(months)
    for candidate in _iter_detail_candidates(bills_df, "bill", "bill_date", "total_amount", valid_months, fx_rate_default) or []:
        row_key = _classify_candidate(candidate["text_map"], candidate["source_type"], row_rules)
        if row_key:
            cogs_rows[row_key][candidate["month"]] += candidate["amount_inr"]

    for candidate in _iter_detail_candidates(journals_df, "journal", "journal_date", "total_amount", valid_months, fx_rate_default) or []:
        row_key = _classify_candidate(candidate["text_map"], candidate["source_type"], row_rules)
        if row_key:
            cogs_rows[row_key][candidate["month"]] += candidate["amount_inr"]

    ramki_monthly_inr = _to_float(company_rules.get("ramki_monthly_usd"), default=12_500) * fx_rate_default
    tech_cost_inr = _to_float(company_rules.get("monthly_technology_cost_inr"), default=25_800)
    techm_start = _to_string(company_rules.get("techm_cogs_start")) or "2025-09"
    dorothea_start = _to_string(company_rules.get("dorothea_cogs_start")) or "2025-10"

    for month_key in months:
        if month_key >= techm_start:
            cogs_rows["ramki"][month_key] = _round_currency(ramki_monthly_inr * 0.5)
        if month_key >= dorothea_start:
            cogs_rows["dorothea"][month_key] = _round_currency(cogs_rows["dorothea"][month_key])
        cogs_rows["technology_costs"][month_key] = _round_currency(tech_cost_inr)

    for row_months in cogs_rows.values():
        for month_key, value in list(row_months.items()):
            row_months[month_key] = _round_currency(value)

    return cogs_rows


def _build_monthly_line_items(
    months: list[str],
    mapped_totals: dict[str, dict[str, float]],
    cogs_rows: dict[str, dict[str, float]],
    company_rules: dict[str, Any],
) -> dict[str, dict[str, float]]:
    monthly_items = {
        "techm_billings": mapped_totals.get("techm_billings", _zero_month_map(months)),
        "bsm_revenue": mapped_totals.get("bsm_revenue", _zero_month_map(months)),
        "fd_interest": mapped_totals.get("fd_interest", _zero_month_map(months)),
        "offshore_cogs": _zero_month_map(months),
        "onshore_cogs": _zero_month_map(months),
        "technology_costs": _zero_month_map(months),
        "ramki_sm": _zero_month_map(months),
        "advertising_marketing": mapped_totals.get("advertising_marketing", _zero_month_map(months)),
        "travel_expenses": mapped_totals.get("travel_expenses", _zero_month_map(months)),
        "meals_entertainment": mapped_totals.get("meals_entertainment", _zero_month_map(months)),
        "rd_salaries": mapped_totals.get("rd_salaries", _zero_month_map(months)),
        "consultant_expense": mapped_totals.get("consultant_expense", _zero_month_map(months)),
        "software_subscriptions": mapped_totals.get("software_subscriptions", _zero_month_map(months)),
        "rent": mapped_totals.get("rent", _zero_month_map(months)),
        "it_internet": mapped_totals.get("it_internet", _zero_month_map(months)),
        "legal": mapped_totals.get("legal", _zero_month_map(months)),
        "audit_non_operating": mapped_totals.get("audit_non_operating", _zero_month_map(months)),
        "other_ga": _zero_month_map(months),
        "cogs_total": _zero_month_map(months),
    }

    ramki_monthly_inr = _to_float(company_rules.get("ramki_monthly_usd"), default=12_500) * _to_float(company_rules.get("fx_rate_inr_per_usd"), default=FX_RATE_DEFAULT)
    techm_start = _to_string(company_rules.get("techm_cogs_start")) or "2025-09"

    for month_key in months:
        monthly_items["ramki_sm"][month_key] = _round_currency(ramki_monthly_inr if month_key < techm_start else ramki_monthly_inr * 0.5)
        monthly_items["offshore_cogs"][month_key] = _round_currency(
            sum(
                cogs_rows[row_key].get(month_key, 0)
                for row_key in ("parashar", "sangeeth", "jitin", "dipti", "dhanashri", "pinkesh", "bsm_crew", "conam_tech", "external_vendors")
            )
        )
        monthly_items["onshore_cogs"][month_key] = _round_currency(
            cogs_rows["ramki"].get(month_key, 0) + cogs_rows["dorothea"].get(month_key, 0)
        )
        monthly_items["technology_costs"][month_key] = _round_currency(cogs_rows["technology_costs"].get(month_key, 0))
        monthly_items["other_ga"][month_key] = _round_currency(
            mapped_totals.get("delivery_india", _zero_month_map(months)).get(month_key, 0)
            + mapped_totals.get("bsm_delivery_contractors", _zero_month_map(months)).get(month_key, 0)
            + mapped_totals.get("external_delivery_vendors", _zero_month_map(months)).get(month_key, 0)
            - sum(
                cogs_rows[row_key].get(month_key, 0)
                for row_key in ("parashar", "sangeeth", "jitin", "dipti", "dhanashri", "pinkesh", "bsm_crew", "conam_tech", "external_vendors")
            )
        )
        monthly_items["cogs_total"][month_key] = _round_currency(
            monthly_items["offshore_cogs"][month_key]
            + monthly_items["onshore_cogs"][month_key]
            + monthly_items["technology_costs"][month_key]
        )

    for month_key in months:
        if monthly_items["other_ga"][month_key] < 0:
            monthly_items["other_ga"][month_key] = 0.0

    return monthly_items


def _build_monthly_preview_dataframe(months: list[str], month_labels: dict[str, str], monthly_items: dict[str, dict[str, float]]) -> pd.DataFrame:
    preview_rows = []
    for label, item_key in MONTHLY_PREVIEW_ITEMS:
        row = {"Line Item": label}
        for month_key in months:
            row[month_labels[month_key]] = _round_currency(monthly_items[item_key][month_key])
        preview_rows.append(row)

    return pd.DataFrame(preview_rows)


def _quarter_formula_from_monthly(month_columns: dict[str, int], target_months: list[str], row_number: int) -> str:
    start_letter = get_column_letter(month_columns[target_months[0]])
    end_letter = get_column_letter(month_columns[target_months[-1]])
    return f"=SUM('{MONTHLY_TEMPLATE_SHEET}'!{start_letter}{row_number}:{end_letter}{row_number})"


def _quarter_formula_from_cogs(cogs_columns: dict[str, int], target_months: list[str], start_row: int, end_row: int | None = None) -> str:
    end_row = end_row or start_row
    start_letter = get_column_letter(cogs_columns[target_months[0]])
    end_letter = get_column_letter(cogs_columns[target_months[-1]])
    return f"=SUM('{COGS_TEMPLATE_SHEET}'!{start_letter}{start_row}:{end_letter}{end_row})"


def _quarter_formula_from_cogs_ranges(
    cogs_columns: dict[str, int],
    target_months: list[str],
    row_ranges: list[tuple[int, int]],
) -> str:
    """Build a quarter formula that adds multiple COGS row ranges."""
    start_letter = get_column_letter(cogs_columns[target_months[0]])
    end_letter = get_column_letter(cogs_columns[target_months[-1]])
    range_terms = [
        f"'{COGS_TEMPLATE_SHEET}'!{start_letter}{start_row}:{end_letter}{end_row}"
        for start_row, end_row in row_ranges
    ]
    return f"=SUM({','.join(range_terms)})"


def _clear_range_values(worksheet, start_row: int, end_row: int, start_column: int, end_column: int) -> None:
    """Clear a rectangular cell range without touching styles."""
    for row_number in range(start_row, end_row + 1):
        for column_number in range(start_column, end_column + 1):
            worksheet.cell(row=row_number, column=column_number).value = None


def _set_column_visibility(worksheet, start_column: int, total_slots: int, visible_slots: int) -> None:
    """Show only the required period columns."""
    for slot_index in range(total_slots):
        column_letter = get_column_letter(start_column + slot_index)
        worksheet.column_dimensions[column_letter].hidden = slot_index >= visible_slots


def _sum_formula(column_letters: list[str], row_number: int) -> str:
    """Build a row sum formula for one or more visible period columns."""
    if len(column_letters) == 1:
        return f"={column_letters[0]}{row_number}"
    return f"=SUM({column_letters[0]}{row_number}:{column_letters[-1]}{row_number})"


def _copy_row_format(worksheet, source_row: int, target_row: int, start_column: int = 2, end_column: int = 14) -> None:
    """Copy formatting from one workbook row to another."""
    for column_number in range(start_column, end_column + 1):
        source_cell = worksheet.cell(source_row, column_number)
        target_cell = worksheet.cell(target_row, column_number)
        target_cell._style = copy(source_cell._style)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height


def _ensure_monthly_sheet_layout(monthly_ws) -> None:
    """Insert extra COGS detail rows once so the monthly sheet can show grouped COGS output."""
    if monthly_ws.max_row >= 34 and monthly_ws["B10"].value == "Offshore COGS":
        return

    monthly_ws.insert_rows(11, amount=3)
    for row_number in (11, 12, 13):
        _copy_row_format(monthly_ws, 10, row_number)


def _apply_report_labels(monthly_ws, quarterly_ws) -> None:
    """Write the requested output labels into the workbook."""
    for row_number, label in MONTHLY_ROW_LABELS.items():
        monthly_ws.cell(row=row_number, column=2).value = label

    for row_number, label in QUARTERLY_ROW_LABELS.items():
        quarterly_ws.cell(row=row_number, column=2).value = label

    for row_number in (12, 13, 14):
        monthly_ws.row_dimensions[row_number].hidden = False

    for row_number in (14, 15, 16):
        quarterly_ws.row_dimensions[row_number].hidden = True


def _set_template_titles(
    quarterly_ws,
    monthly_ws,
    cogs_ws,
    report_period: dict[str, Any],
    fx_rate: float,
    company_rules: dict[str, Any],
) -> None:
    quarter_headers = report_period["summary_headers"]
    months = report_period["months"]

    quarterly_ws["B1"] = f"MIDOFFICE DATA  |  Profit & Loss Statement  |  {report_period['header_title']}"
    quarterly_ws["B2"] = (
        "India Parent (Midoffice Data Solutions Pvt Ltd) + US Subsidiary (Midoffice Data International Inc.)"
        f"  |  Accrual Basis  |  FX: ₹{int(fx_rate)}/USD  |  COGS: {_format_month_key(_to_string(company_rules.get('techm_cogs_start')) or '2025-09')} onwards  |  All amounts in INR (₹)"
    )

    monthly_ws["B1"] = (
        f"Monthly P&L Detail  |  {report_period['header_title']}  |  Automated from Zoho + BigQuery"
    )
    cogs_ws["B1"] = (
        f"COGS Allocation Detail  |  {report_period['header_title']}  |  "
        f"TechM COGS from {_format_month_key(_to_string(company_rules.get('techm_cogs_start')) or '2025-09')}  |  "
        f"BSM delivery from {_format_month_key(_to_string(company_rules.get('bsm_delivery_cogs_start')) or '2025-07')}"
    )

    for column_offset in range(3, 7):
        quarterly_ws.cell(4, column_offset).value = None
    for column_offset, header_label in enumerate(quarter_headers, start=3):
        quarterly_ws.cell(4, column_offset).value = header_label

    quarterly_ws["G4"] = report_period["total_column_label"]
    quarterly_ws["H4"] = "% Rev"

    for column_number in range(3, 15):
        monthly_ws.cell(3, column_number).value = None
    for column_number, month_period in enumerate(months, start=3):
        monthly_ws.cell(3, column_number).value = month_period["label"]

    for column_number in range(5, 17):
        cogs_ws.cell(3, column_number).value = None
    for column_number, month_period in enumerate(months, start=5):
        cogs_ws.cell(3, column_number).value = month_period["label"]

    _set_column_visibility(monthly_ws, 3, 12, len(months))
    _set_column_visibility(cogs_ws, 5, 12, len(months))
    _set_column_visibility(quarterly_ws, 3, 4, len(quarter_headers))


def _populate_monthly_sheet(monthly_ws, months: list[str], monthly_items: dict[str, dict[str, float]], cogs_rows: dict[str, dict[str, float]]) -> None:
    month_columns = _month_column_map(3, months)
    row_map = {
        "techm_billings": 5,
        "bsm_revenue": 6,
        "fd_interest": 7,
        "offshore_cogs": 10,
        "onshore_cogs": 11,
        "technology_costs": 12,
        "ramki_sm": 16,
        "advertising_marketing": 17,
        "travel_expenses": 18,
        "meals_entertainment": 19,
        "rd_salaries": 22,
        "consultant_expense": 23,
        "software_subscriptions": 24,
        "rent": 27,
        "it_internet": 28,
        "legal": 29,
        "audit_non_operating": 30,
        "other_ga": 31,
    }

    for item_key, row_number in row_map.items():
        for month_key, column_number in month_columns.items():
            monthly_ws.cell(row=row_number, column=column_number).value = _round_currency(monthly_items[item_key][month_key])

    for month_key, column_number in month_columns.items():
        current_letter = get_column_letter(column_number)
        monthly_ws[f"{current_letter}8"] = f"=SUM({current_letter}5:{current_letter}7)"
        monthly_ws[f"{current_letter}13"] = f"=SUM({current_letter}10:{current_letter}12)"
        monthly_ws[f"{current_letter}14"] = f"={current_letter}8-{current_letter}13"
        monthly_ws[f"{current_letter}20"] = f"=SUM({current_letter}16:{current_letter}19)"
        monthly_ws[f"{current_letter}25"] = f"=SUM({current_letter}22:{current_letter}24)"
        monthly_ws[f"{current_letter}32"] = f"=SUM({current_letter}27:{current_letter}31)"
        monthly_ws[f"{current_letter}33"] = f"=SUM({current_letter}20,{current_letter}25,{current_letter}32)"
        monthly_ws[f"{current_letter}34"] = f"={current_letter}14-{current_letter}33"


def _populate_cogs_sheet(cogs_ws, months: list[str], cogs_rows: dict[str, dict[str, float]]) -> None:
    month_columns = _month_column_map(5, months)
    row_map = {
        "parashar": 4,
        "sangeeth": 5,
        "jitin": 6,
        "dipti": 7,
        "dhanashri": 8,
        "pinkesh": 9,
        "ramki": 10,
        "dorothea": 11,
        "bsm_crew": 12,
        "conam_tech": 13,
        "external_vendors": 14,
    }

    last_month_letter = get_column_letter(month_columns[months[-1]])
    for row_key, row_number in row_map.items():
        for month_key, column_number in month_columns.items():
            cogs_ws.cell(row=row_number, column=column_number).value = _round_currency(cogs_rows[row_key][month_key])
        first_letter = get_column_letter(month_columns[months[0]])
        cogs_ws[f"Q{row_number}"] = f"=SUM({first_letter}{row_number}:{last_month_letter}{row_number})"

    cogs_ws["D15"] = None
    for month_key, column_number in month_columns.items():
        current_letter = get_column_letter(column_number)
        cogs_ws[f"{current_letter}15"] = f"=SUM({current_letter}4:{current_letter}14)"
    cogs_ws["Q15"] = f"=SUM(E15:{last_month_letter}15)"


def _populate_quarterly_sheet(
    quarterly_ws,
    months: list[str],
    company_rules: dict[str, Any],
    report_period: dict[str, Any],
) -> None:
    month_columns = _month_column_map(3, months)
    cogs_columns = _month_column_map(5, months)
    summary_periods = report_period["summary_periods"]
    summary_letters = [get_column_letter(column_number) for column_number in range(3, 3 + len(summary_periods))]

    monthly_row_map = {
        6: 5,
        7: 6,
        8: 7,
        25: 16,
        26: 17,
        27: 18,
        28: 19,
        32: 22,
        33: 23,
        34: 24,
        38: 27,
        39: 28,
        40: 29,
        41: 30,
        42: 31,
    }
    percent_rows = [6, 7, 8, 9, 12, 13, 17, 18, 20, 25, 26, 27, 28, 29, 32, 33, 34, 35, 38, 39, 40, 41, 42, 43, 45, 47]
    tech_cost_per_month = _to_float(company_rules.get("monthly_technology_cost_inr"), default=25_800)

    for quarter_index, period_bucket in enumerate(summary_periods, start=3):
        quarter_letter = get_column_letter(quarter_index)
        quarter_month_keys = period_bucket["month_keys"]
        for quarterly_row, monthly_row in monthly_row_map.items():
            quarterly_ws[f"{quarter_letter}{quarterly_row}"] = _quarter_formula_from_monthly(month_columns, quarter_month_keys, monthly_row)

        quarterly_ws[f"{quarter_letter}12"] = _quarter_formula_from_cogs_ranges(
            cogs_columns,
            quarter_month_keys,
            [(4, 9), (12, 14)],
        )
        quarterly_ws[f"{quarter_letter}13"] = _quarter_formula_from_cogs_ranges(
            cogs_columns,
            quarter_month_keys,
            [(10, 11)],
        )
        quarterly_ws[f"{quarter_letter}14"] = None
        quarterly_ws[f"{quarter_letter}15"] = None
        quarterly_ws[f"{quarter_letter}16"] = None

        quarterly_ws[f"{quarter_letter}17"] = f"={tech_cost_per_month}*{len(quarter_month_keys)}"
        quarterly_ws[f"{quarter_letter}9"] = f"=SUM({quarter_letter}6:{quarter_letter}8)"
        quarterly_ws[f"{quarter_letter}18"] = f"=SUM({quarter_letter}12:{quarter_letter}17)"
        quarterly_ws[f"{quarter_letter}20"] = f"={quarter_letter}9-{quarter_letter}18"
        quarterly_ws[f"{quarter_letter}21"] = f"=IFERROR({quarter_letter}20/{quarter_letter}9,0)"
        quarterly_ws[f"{quarter_letter}29"] = f"=SUM({quarter_letter}25:{quarter_letter}28)"
        quarterly_ws[f"{quarter_letter}35"] = f"=SUM({quarter_letter}32:{quarter_letter}34)"
        quarterly_ws[f"{quarter_letter}43"] = f"=SUM({quarter_letter}38:{quarter_letter}42)"
        quarterly_ws[f"{quarter_letter}45"] = f"=SUM({quarter_letter}29,{quarter_letter}35,{quarter_letter}43)"
        quarterly_ws[f"{quarter_letter}47"] = f"={quarter_letter}20-{quarter_letter}45"
        quarterly_ws[f"{quarter_letter}48"] = f"=IFERROR({quarter_letter}47/{quarter_letter}9,0)"

    for row_number in [6, 7, 8, 12, 13, 17, 25, 26, 27, 28, 32, 33, 34, 38, 39, 40, 41, 42]:
        quarterly_ws[f"G{row_number}"] = _sum_formula(summary_letters, row_number)

    quarterly_ws["G14"] = None
    quarterly_ws["G15"] = None
    quarterly_ws["G16"] = None
    quarterly_ws["G9"] = "=SUM(G6:G8)"
    quarterly_ws["G18"] = "=SUM(G12:G17)"
    quarterly_ws["G20"] = "=G9-G18"
    quarterly_ws["G21"] = "=IFERROR(G20/G9,0)"
    quarterly_ws["G29"] = "=SUM(G25:G28)"
    quarterly_ws["G35"] = "=SUM(G32:G34)"
    quarterly_ws["G43"] = "=SUM(G38:G42)"
    quarterly_ws["G45"] = "=SUM(G29,G35,G43)"
    quarterly_ws["G47"] = "=G20-G45"
    quarterly_ws["G48"] = "=IFERROR(G47/G9,0)"

    for row_number in percent_rows:
        quarterly_ws[f"H{row_number}"] = f"=IFERROR(G{row_number}/$G$9,0)"
    quarterly_ws["H14"] = None
    quarterly_ws["H15"] = None
    quarterly_ws["H16"] = None
    quarterly_ws["H21"] = None
    quarterly_ws["H48"] = None


def _populate_key_metrics_and_notes(
    quarterly_ws,
    monthly_ws,
    months: list[str],
    dashboard_summary_df: pd.DataFrame,
    note_totals: dict[str, dict[str, float]],
    cogs_rows: dict[str, dict[str, float]],
    company_rules: dict[str, Any],
    report_period: dict[str, Any],
    invoices_df: pd.DataFrame,
    bills_df: pd.DataFrame,
    journals_df: pd.DataFrame,
) -> None:
    total_label = "Gross Margin (FY)" if report_period["period_type"] == "full_year" else "Gross Margin (Period)"
    quarterly_ws["B52"] = total_label
    quarterly_ws["B53"] = "Peak Quarter Gross Margin"
    quarterly_ws["B54"] = "Monthly OpEx Run-Rate"
    quarterly_ws["B55"] = "Monthly Gross Burn (run-rate)"
    quarterly_ws["B56"] = "Monthly Net (run-rate)"
    quarterly_ws["B57"] = "Revenue per Delivery FTE"
    quarterly_ws["B58"] = "Total People Cost (FY)"

    valid_run_rate_months = months[-1:] if len(months) == 1 else months[-min(3, len(months)) :]
    monthly_columns = _month_column_map(3, months)
    run_rate_start = get_column_letter(monthly_columns[valid_run_rate_months[0]])
    run_rate_end = get_column_letter(monthly_columns[valid_run_rate_months[-1]])
    last_summary_letter = get_column_letter(2 + len(report_period["summary_periods"]))
    average_bonus_adjustment = _round_currency(
        sum(note_totals.get("bonus", {}).get(month_key, 0) for month_key in valid_run_rate_months)
        / max(len(valid_run_rate_months), 1)
    )

    quarterly_ws["C52"] = '=TEXT(G21,"0.0%")'
    quarterly_ws["D52"] = "Gross margin based on automated revenue and COGS roll-ups for the selected report."
    quarterly_ws["C53"] = f'=TEXT(MAX(C21:{last_summary_letter}21),"0.0%")'
    quarterly_ws["D53"] = "Highest visible summary-period gross margin from the generated quarterly view."
    quarterly_ws["C54"] = (
        f'=TEXT(AVERAGE(\'{MONTHLY_TEMPLATE_SHEET}\'!{run_rate_start}33:{run_rate_end}33)-{average_bonus_adjustment},"₹#,##0")'
    )
    quarterly_ws["D54"] = f"Run-rate uses {', '.join(pd.Timestamp(f'{month_key}-01').strftime('%b %Y') for month_key in valid_run_rate_months)}."
    quarterly_ws["C55"] = (
        f'=TEXT(AVERAGE(\'{MONTHLY_TEMPLATE_SHEET}\'!{run_rate_start}13:{run_rate_end}13)'
        f'+AVERAGE(\'{MONTHLY_TEMPLATE_SHEET}\'!{run_rate_start}33:{run_rate_end}33)-{average_bonus_adjustment},"₹#,##0")'
    )
    quarterly_ws["D55"] = "Gross burn is average monthly COGS plus run-rate operating expense, excluding bonus items."
    quarterly_ws["C56"] = f'=TEXT(AVERAGE(\'{MONTHLY_TEMPLATE_SHEET}\'!{run_rate_start}34:{run_rate_end}34)+{average_bonus_adjustment},"₹#,##0")'
    quarterly_ws["D56"] = "Monthly net run-rate adds back configured bonus items from the chosen run-rate months."
    quarterly_ws["C57"] = '=TEXT(IFERROR(G9/MAX(COUNTA(\'COGS Allocation Working\'!B4:B9),1),0),"₹#,##0")'
    quarterly_ws["D57"] = "Simple revenue-per-delivery-head view using the named delivery team rows in the allocation sheet."
    quarterly_ws["C58"] = '=TEXT(SUM(\'COGS Allocation Working\'!Q4:Q12,G25,G32),"₹#,##0")'
    quarterly_ws["D58"] = "People cost view includes delivery people cost, Onshore Consultant Experience, and R&D salaries inside the selected period."
    quarterly_ws["C59"] = "-"
    quarterly_ws["D59"] = "Closing cash is left blank until balance-sheet or bank balance data is added to the MIS pipeline."
    quarterly_ws["C60"] = "-"
    quarterly_ws["D60"] = "Cash runway is left blank until cash balances are sourced automatically."

    bonus_total = sum(note_totals.get("bonus", {}).values())
    one_time_total = sum(note_totals.get("one_time", {}).values())
    bsm_provision_total = sum(note_totals.get("bsm_provision", {}).values())
    techm_receivable_total = sum(note_totals.get("techm_receivable", {}).values())
    invoice_count = len(invoices_df.index)
    bill_count = len(bills_df.index)
    journal_count = len(journals_df.index)
    ramki_total = sum(cogs_rows["ramki"].values()) + sum(
        monthly_ws.cell(16, column_number).value or 0
        for column_number in range(3, 3 + len(months))
    )

    quarterly_ws["B62"] = (
        f"⚑ Bonus review: {_format_inr_short(bonus_total)} identified against the MIS note-tracking rules. "
        "One-time bonus items remain visible in the P&L and are excluded only from run-rate metrics."
    )
    quarterly_ws["B63"] = (
        f"⚑ One-time expense review: {_format_inr_short(one_time_total)} flagged from source descriptions and notes."
    )
    quarterly_ws["B64"] = (
        f"★ BSM provision watchlist: {_format_inr_short(bsm_provision_total)} flagged through configured MIS keyword rules."
    )
    quarterly_ws["B65"] = (
        f"     TechM receivable watchlist: {_format_inr_short(techm_receivable_total)} tagged from invoice metadata for follow-up."
    )
    quarterly_ws["B66"] = (
        f"     Onshore Consultant Experience uses the Ramki assumption of $12,500/month with 100% S&M Apr-Aug and 50% COGS / 50% S&M from Sep onward. "
        f"Generated report-period Ramki cost: {_format_inr_short(ramki_total)}."
    )
    quarterly_ws["B67"] = (
        f"     Workbook generated from Zoho/BigQuery detail rows for {report_period['period_name']}. Source counts included in this run: "
        f"{invoice_count} invoices, {bill_count} bills, {journal_count} journals."
    )


def _build_metrics(
    financial_year: str,
    report_period: dict[str, Any],
    monthly_sheet_preview: pd.DataFrame,
    dashboard_summary_df: pd.DataFrame,
    invoices_df: pd.DataFrame,
    bills_df: pd.DataFrame,
    journals_df: pd.DataFrame,
) -> dict[str, str]:
    revenue_rows = monthly_sheet_preview[
        monthly_sheet_preview["Line Item"].isin(
            ["Client Revenue - Professional Services", "Client Revenue - Product", "Other Income"]
        )
    ]
    expense_rows = monthly_sheet_preview[
        monthly_sheet_preview["Line Item"].isin(
            [
                "COGS Total",
                "Onshore Consultant Experience",
                "Advertising & Marketing",
                "Travel Expenses",
                "Meals & Entertainment",
                "R&D Salaries",
                "Consultant & Contractor Expense",
                "Software Subs",
                "Office Rent",
                "IT & Internet",
                "Legal",
                "Audit & Non-Op",
                "Other G&A",
            ]
        )
    ]
    numeric_columns = [column_name for column_name in monthly_sheet_preview.columns if column_name != "Line Item"]

    revenue_total = float(revenue_rows[numeric_columns].sum().sum()) if not revenue_rows.empty else 0.0
    expense_total = float(expense_rows[numeric_columns].sum().sum()) if not expense_rows.empty else 0.0
    profit_total = revenue_total - expense_total

    return {
        "Financial Year": financial_year,
        "Report Period": report_period["period_name"],
        "Organization": _first_value(dashboard_summary_df, "source_org_name", "All Organizations"),
        "Reporting Currency": _first_value(dashboard_summary_df, "reporting_currency", "INR"),
        "Revenue": _format_currency(revenue_total),
        "Expenses": _format_currency(expense_total),
        "Profit": _format_currency(profit_total),
        "Journal Adjustments": _format_currency(0),
        "Invoices": str(len(invoices_df.index)),
        "Bills": str(len(bills_df.index)),
        "Contacts": str(int(_first_value(dashboard_summary_df, "contact_count"))),
        "Journals": str(len(journals_df.index)),
    }


def get_mis_metrics(
    financial_year: int | str,
    monthly_pl_df: pd.DataFrame | None = None,
    dashboard_summary_df: pd.DataFrame | None = None,
    project_id: str | None = None,
    org_filter: str | None = None,
    period_type: str = "full_year",
    selected_month: int | None = None,
    selected_quarter: str | None = None,
    selected_half: str | None = None,
    custom_start_date: date | datetime | None = None,
    custom_end_date: date | datetime | None = None,
) -> dict[str, str]:
    """Return Streamlit metric-card values from Consume layer data."""
    financial_year_start = parse_financial_year_start(financial_year)
    report_period = get_selected_report_period(
        financial_year_start,
        period_type,
        selected_month=selected_month,
        selected_quarter=selected_quarter,
        selected_half=selected_half,
        custom_start_date=custom_start_date,
        custom_end_date=custom_end_date,
    )
    if monthly_pl_df is None or dashboard_summary_df is None:
        monthly_pl_df, dashboard_summary_df = fetch_gold_mis_data(project_id, org_filter=org_filter, report_period=report_period)

    monthly_pl_df = monthly_pl_df if monthly_pl_df is not None else pd.DataFrame()
    monthly_pl_df = _filter_dataframe_by_date_range(
        monthly_pl_df,
        "report_month",
        report_period["start_date"],
        report_period["end_date"],
    )
    dashboard_summary_df = dashboard_summary_df if dashboard_summary_df is not None else pd.DataFrame()

    revenue_amount = _safe_sum(monthly_pl_df, "revenue_amount")
    expense_amount = _safe_sum(monthly_pl_df, "expense_amount")
    journal_adjustment_amount = _safe_sum(monthly_pl_df, "journal_adjustment_amount")
    profit_amount = _safe_sum(monthly_pl_df, "profit_amount") or (revenue_amount - expense_amount + journal_adjustment_amount)

    return {
        "Financial Year": report_period["fy_label"],
        "Report Period": report_period["period_name"],
        "Organization": _first_value(dashboard_summary_df, "source_org_name", "All Organizations"),
        "Reporting Currency": _first_value(dashboard_summary_df, "reporting_currency", "INR"),
        "Revenue": _format_currency(revenue_amount),
        "Expenses": _format_currency(expense_amount),
        "Profit": _format_currency(profit_amount),
        "Journal Adjustments": _format_currency(journal_adjustment_amount),
        "Invoices": str(int(_first_value(dashboard_summary_df, "invoice_count"))),
        "Bills": str(int(_first_value(dashboard_summary_df, "bill_count"))),
        "Contacts": str(int(_first_value(dashboard_summary_df, "contact_count"))),
    }


def _prepare_dashboard_kpis(metrics: dict[str, str]) -> pd.DataFrame:
    return pd.DataFrame([{"KPI": key, "Value": value} for key, value in metrics.items()])


def _generate_template_workbook(
    template_path: Path,
    destination_path: Path,
    report_period: dict[str, Any],
    months: list[str],
    mapped_totals: dict[str, dict[str, float]],
    cogs_rows: dict[str, dict[str, float]],
    dashboard_summary_df: pd.DataFrame,
    company_rules: dict[str, Any],
    note_totals: dict[str, dict[str, float]],
    invoices_df: pd.DataFrame,
    bills_df: pd.DataFrame,
    journals_df: pd.DataFrame,
) -> None:
    workbook = load_workbook(template_path)
    quarterly_ws = workbook[QUARTERLY_TEMPLATE_SHEET]
    monthly_ws = workbook[MONTHLY_TEMPLATE_SHEET]
    cogs_ws = workbook[COGS_TEMPLATE_SHEET]
    _ensure_monthly_sheet_layout(monthly_ws)

    fx_rate = _to_float(company_rules.get("fx_rate_inr_per_usd"), default=FX_RATE_DEFAULT)
    monthly_items = _build_monthly_line_items(months, mapped_totals, cogs_rows, company_rules)

    _clear_range_values(quarterly_ws, 4, 48, 3, 8)
    _clear_range_values(monthly_ws, 3, 34, 3, 14)
    _clear_range_values(cogs_ws, 3, 15, 5, 17)

    _set_template_titles(quarterly_ws, monthly_ws, cogs_ws, report_period, fx_rate, company_rules)
    _apply_report_labels(monthly_ws, quarterly_ws)
    _populate_cogs_sheet(cogs_ws, months, cogs_rows)
    _populate_monthly_sheet(monthly_ws, months, monthly_items, cogs_rows)
    _populate_quarterly_sheet(quarterly_ws, months, company_rules, report_period)
    _populate_key_metrics_and_notes(
        quarterly_ws,
        monthly_ws,
        months,
        dashboard_summary_df,
        note_totals,
        cogs_rows,
        company_rules,
        report_period,
        invoices_df,
        bills_df,
        journals_df,
    )

    quarterly_ws.freeze_panes = "C5"
    monthly_ws.freeze_panes = "C4"
    cogs_ws.freeze_panes = "D4"
    workbook._sheets = [quarterly_ws, monthly_ws, cogs_ws]
    workbook.calculation.forceFullCalc = True
    workbook.calculation.fullCalcOnLoad = True
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination_path)


def generate_mis_report(
    financial_year: int | str,
    output_dir: str | Path,
    project_id: str | None = None,
    org_filter: str | None = None,
    period_type: str = "full_year",
    selected_month: int | None = None,
    selected_quarter: str | None = None,
    selected_half: str | None = None,
    custom_start_date: date | datetime | None = None,
    custom_end_date: date | datetime | None = None,
    monthly_pl_df: pd.DataFrame | None = None,
    dashboard_summary_df: pd.DataFrame | None = None,
    bills_df: pd.DataFrame | None = None,
    invoices_df: pd.DataFrame | None = None,
    journals_df: pd.DataFrame | None = None,
) -> dict[str, Any]:
    """Generate the CEO-format MIS workbook from detailed Zoho/BigQuery data."""
    financial_year_start = parse_financial_year_start(financial_year)
    report_period = get_selected_report_period(
        financial_year_start,
        period_type,
        selected_month=selected_month,
        selected_quarter=selected_quarter,
        selected_half=selected_half,
        custom_start_date=custom_start_date,
        custom_end_date=custom_end_date,
    )
    template_path = _template_path()
    if not template_path.exists():
        raise FileNotFoundError(f"MIS template workbook not found at {template_path}")

    mapping_config = _read_yaml(_mis_mapping_path()).get("mis_report", {})
    cogs_config = _read_yaml(_cogs_config_path())
    company_rules = dict(cogs_config.get("company_rules", {}))
    company_rules["fx_rate_inr_per_usd"] = _to_float(cogs_config.get("fx_rate_inr_per_usd"), default=FX_RATE_DEFAULT)
    months = report_period["month_keys"]

    monthly_pl_df = monthly_pl_df if monthly_pl_df is not None else pd.DataFrame()
    dashboard_summary_df = dashboard_summary_df if dashboard_summary_df is not None else pd.DataFrame()
    if monthly_pl_df.empty and dashboard_summary_df.empty:
        try:
            monthly_pl_df, dashboard_summary_df = fetch_gold_mis_data(project_id, org_filter=org_filter, report_period=report_period)
        except Exception:
            monthly_pl_df = pd.DataFrame()
            dashboard_summary_df = pd.DataFrame()
    monthly_pl_df = _filter_dataframe_by_date_range(
        monthly_pl_df,
        "report_month",
        report_period["start_date"],
        report_period["end_date"],
    )

    if invoices_df is None or bills_df is None or journals_df is None:
        fetched_invoices_df, fetched_bills_df, fetched_journals_df = fetch_detailed_mis_data(
            project_id,
            org_filter=org_filter,
            start_date=report_period["start_date"],
            end_date=report_period["end_date"],
        )
        invoices_df = fetched_invoices_df if invoices_df is None else invoices_df
        bills_df = fetched_bills_df if bills_df is None else bills_df
        journals_df = fetched_journals_df if journals_df is None else journals_df

    invoices_df = invoices_df if invoices_df is not None else pd.DataFrame()
    bills_df = bills_df if bills_df is not None else pd.DataFrame()
    journals_df = journals_df if journals_df is not None else pd.DataFrame()
    invoices_df = _filter_dataframe_by_date_range(invoices_df, "invoice_date", report_period["start_date"], report_period["end_date"])
    bills_df = _filter_dataframe_by_date_range(bills_df, "bill_date", report_period["start_date"], report_period["end_date"])
    journals_df = _filter_dataframe_by_date_range(journals_df, "journal_date", report_period["start_date"], report_period["end_date"])

    fx_rate_default = company_rules["fx_rate_inr_per_usd"]
    account_mappings = mapping_config.get("account_mappings", [])
    revenue_rules = _merge_account_mapping_rules(
        mapping_config.get("revenue", {}),
        account_mappings,
        "revenue",
        org_filter,
    )
    expense_rules = _merge_account_mapping_rules(
        mapping_config.get("expense", {}),
        account_mappings,
        "expense",
        org_filter,
    )
    note_rules = mapping_config.get("note_tracking", {})

    mapped_totals = _aggregate_rule_totals(months, fx_rate_default, invoices_df, bills_df, journals_df, {**revenue_rules, **expense_rules})
    note_totals = _aggregate_rule_totals(months, fx_rate_default, invoices_df, bills_df, journals_df, note_rules)
    cogs_rows = _aggregate_cogs_detail(months, fx_rate_default, bills_df, journals_df, company_rules)
    monthly_items = _build_monthly_line_items(months, mapped_totals, cogs_rows, company_rules)
    monthly_preview_df = _build_monthly_preview_dataframe(months, _month_labels(months), monthly_items)

    destination_folder = Path(output_dir)
    destination_folder.mkdir(parents=True, exist_ok=True)
    report_path = destination_folder / f"MIS_PL_{_financial_year_token(financial_year_start)}_{report_period['file_suffix']}.xlsx"

    _generate_template_workbook(
        template_path=template_path,
        destination_path=report_path,
        report_period=report_period,
        months=months,
        mapped_totals=mapped_totals,
        cogs_rows=cogs_rows,
        dashboard_summary_df=dashboard_summary_df,
        company_rules=company_rules,
        note_totals=note_totals,
        invoices_df=invoices_df,
        bills_df=bills_df,
        journals_df=journals_df,
    )

    metrics = _build_metrics(
        report_period["fy_label"],
        report_period,
        monthly_preview_df,
        dashboard_summary_df,
        invoices_df,
        bills_df,
        journals_df,
    )
    return {
        "status": "success",
        "message": f"MIS report generated for {report_period['period_name']}.",
        "metrics": metrics,
        "report_path": report_path,
        "is_placeholder": False,
        "summary": monthly_preview_df,
        "monthly_preview": monthly_preview_df,
        "report_period": report_period,
        "dashboard_kpis": _prepare_dashboard_kpis(metrics),
    }


def _consolidated_rows(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Use the Consume-layer all-organizations rows when they are available."""
    if dataframe.empty or "source_org_key" not in dataframe.columns:
        return dataframe.copy()
    all_rows = dataframe[dataframe["source_org_key"].fillna("").astype(str).str.lower() == "all"]
    return all_rows.copy() if not all_rows.empty else dataframe.copy()


def _consolidated_pl_values(monthly_pl_df: pd.DataFrame) -> tuple[dict[str, float], list[str]]:
    """Build a conservative consolidated P&L without inventing classifications."""
    working_df = _consolidated_rows(monthly_pl_df)
    revenue = _safe_sum(working_df, "revenue_amount")
    total_expenses = _safe_sum(working_df, "expense_amount")
    cogs = _safe_sum(working_df, "cogs_amount")
    other_income = _safe_sum(working_df, "other_income_amount")
    other_expenses = _safe_sum(working_df, "other_expense_amount")

    journal_adjustments = _numeric_column(working_df, "journal_adjustment_amount")
    if "other_income_amount" not in working_df.columns:
        other_income = float(journal_adjustments.clip(lower=0).sum())
    if "other_expense_amount" not in working_df.columns:
        other_expenses = abs(float(journal_adjustments.clip(upper=0).sum()))

    if "operating_expense_amount" in working_df.columns:
        operating_expenses = _safe_sum(working_df, "operating_expense_amount")
    else:
        operating_expenses = max(total_expenses - cogs - other_expenses, 0.0)

    gross_profit = revenue - cogs
    operating_profit = gross_profit - operating_expenses
    net_profit = operating_profit + other_income - other_expenses
    messages: list[str] = []
    if working_df.empty:
        messages.append("No consolidated P&L data was available for the selected period; the workbook contains a blank-safe report.")
    if "cogs_amount" not in working_df.columns:
        messages.append("COGS classification was unavailable, so available bill expenses are presented under Operating Expenses.")
    if "other_income_amount" not in working_df.columns or "other_expense_amount" not in working_df.columns:
        messages.append("Other income and expense use the sign of available journal adjustments where detailed classification is unavailable.")

    return {
        "Revenue": revenue,
        "COGS": cogs,
        "Gross Profit": gross_profit,
        "Operating Expenses": operating_expenses,
        "Operating Profit / EBITDA": operating_profit,
        "Other Income": other_income,
        "Other Expenses": other_expenses,
        "Net Profit": net_profit,
    }, messages


def fetch_consolidated_balance_sheet_data(
    project_id: str | None,
    as_of_date: date,
) -> pd.DataFrame:
    """Fetch available India/US INR balances at the selected period end date."""
    from google.cloud import bigquery

    resolved_project_id = _get_project_id(project_id)
    client = bigquery.Client(project=resolved_project_id)
    as_of_sql = _date_to_sql(as_of_date)
    query = f"""
        SELECT
            'Receivables' AS line_item,
            SUM(COALESCE(balance_amount_inr, 0)) AS amount_inr,
            COUNT(*) AS source_records
        FROM {_table_name(resolved_project_id, INVOICES_VIEW)}
        WHERE invoice_date <= DATE '{as_of_sql}'
        UNION ALL
        SELECT
            'Payables' AS line_item,
            SUM(COALESCE(balance_amount_inr, 0)) AS amount_inr,
            COUNT(*) AS source_records
        FROM {_table_name(resolved_project_id, BILLS_VIEW)}
        WHERE bill_date <= DATE '{as_of_sql}'
    """
    return _query_to_dataframe(client, query)


def _first_numeric_value(dataframe: pd.DataFrame, column_names: list[str]) -> float | None:
    """Return a summed numeric field when the field exists, preserving missingness."""
    for column_name in column_names:
        if column_name in dataframe.columns:
            numeric = pd.to_numeric(dataframe[column_name], errors="coerce")
            if numeric.notna().any():
                return float(numeric.fillna(0).sum())
    return None


def _balance_sheet_values(balance_sheet_df: pd.DataFrame) -> tuple[dict[str, float | None], list[str]]:
    """Normalize long- or wide-form available balance-sheet inputs."""
    values: dict[str, float | None] = {
        "Bank / Cash": None,
        "Receivables": None,
        "Other Assets": None,
        "Payables": None,
        "Tax Liabilities": None,
        "Equity / Retained Earnings": None,
    }
    aliases = {
        "bankcash": "Bank / Cash",
        "cash": "Bank / Cash",
        "bank": "Bank / Cash",
        "receivables": "Receivables",
        "accountsreceivable": "Receivables",
        "otherassets": "Other Assets",
        "payables": "Payables",
        "accountspayable": "Payables",
        "taxliabilities": "Tax Liabilities",
        "equity": "Equity / Retained Earnings",
        "retainedearnings": "Equity / Retained Earnings",
        "equityretainedearnings": "Equity / Retained Earnings",
    }

    if not balance_sheet_df.empty and "line_item" in balance_sheet_df.columns:
        amount_column = "amount_inr" if "amount_inr" in balance_sheet_df.columns else "amount"
        if amount_column in balance_sheet_df.columns:
            for _, row in balance_sheet_df.iterrows():
                key = re.sub(r"[^a-z]", "", str(row.get("line_item", "")).lower())
                target = aliases.get(key)
                amount = pd.to_numeric(pd.Series([row.get(amount_column)]), errors="coerce").iloc[0]
                if target and not pd.isna(amount):
                    values[target] = float(amount) + float(values[target] or 0)

    wide_columns = {
        "Bank / Cash": ["bank_cash", "cash_amount", "bank_balance_amount"],
        "Receivables": ["receivables", "receivables_amount", "invoice_outstanding_amount"],
        "Other Assets": ["other_assets", "other_assets_amount"],
        "Payables": ["payables", "payables_amount", "bill_outstanding_amount"],
        "Tax Liabilities": ["tax_liabilities", "tax_liability_amount"],
        "Equity / Retained Earnings": ["equity_retained_earnings", "retained_earnings", "equity_amount"],
    }
    for label, column_names in wide_columns.items():
        if values[label] is None:
            values[label] = _first_numeric_value(balance_sheet_df, column_names)

    missing_labels = [label for label, value in values.items() if value is None]
    messages: list[str] = []
    if balance_sheet_df.empty:
        messages.append("No consolidated balance-sheet data was available as of the selected date; the workbook contains a blank-safe report.")
    if missing_labels:
        messages.append("Unavailable categories are shown as blank: " + ", ".join(missing_labels) + ".")
    messages.append("Receivable and payable balances reflect the latest available records dated on or before the as-of date.")
    return values, messages


def _filter_balance_sheet_as_of(balance_sheet_df: pd.DataFrame, as_of_date: date) -> pd.DataFrame:
    """Use the latest supplied balance snapshot on or before the period end."""
    if balance_sheet_df.empty:
        return balance_sheet_df.copy()
    for column_name in ["as_of_date", "balance_date", "report_date"]:
        if column_name not in balance_sheet_df.columns:
            continue
        parsed_dates = pd.to_datetime(balance_sheet_df[column_name], errors="coerce").dt.date
        eligible_df = balance_sheet_df[parsed_dates <= as_of_date].copy()
        if eligible_df.empty:
            return eligible_df
        eligible_dates = pd.to_datetime(eligible_df[column_name], errors="coerce").dt.date
        return eligible_df[eligible_dates == eligible_dates.max()].copy()
    return balance_sheet_df.copy()


def _style_consolidated_statement(worksheet, title: str, subtitle: str) -> None:
    """Apply a compact management-report style to a consolidated statement."""
    dark_blue = "1F2D4E"
    mid_blue = "4472C4"
    light_blue = "D6E4F7"
    thin_gray = Side(style="thin", color="D0D7E2")
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = title
    worksheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=12)
    worksheet["A1"].alignment = Alignment(horizontal="left")
    worksheet.merge_cells("A2:B2")
    worksheet["A2"] = subtitle
    worksheet["A2"].fill = PatternFill("solid", fgColor=mid_blue)
    worksheet["A2"].font = Font(color="FFFFFF", italic=True, size=9)
    worksheet["A4"] = "Line Item"
    worksheet["B4"] = "Amount (INR)"
    for cell in worksheet[4]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 24
    worksheet.freeze_panes = "A5"
    for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
    for cell in worksheet["B"]:
        if cell.row >= 5:
            cell.number_format = '₹#,##0;[Red](₹#,##0);-'
            cell.alignment = Alignment(horizontal="right")
    for row_number in range(5, worksheet.max_row + 1):
        label = str(worksheet.cell(row_number, 1).value or "")
        if label in {"Revenue", "COGS", "Operating Expenses", "Other Income", "Other Expenses", "Assets", "Liabilities", "Equity"}:
            for column_number in (1, 2):
                cell = worksheet.cell(row_number, column_number)
                cell.fill = PatternFill("solid", fgColor=mid_blue)
                cell.font = Font(color="FFFFFF", bold=True)
        elif label.startswith("Total ") or label in {"Gross Profit", "Operating Profit / EBITDA", "Net Profit", "Current Assets"}:
            for column_number in (1, 2):
                cell = worksheet.cell(row_number, column_number)
                cell.fill = PatternFill("solid", fgColor=light_blue)
                cell.font = Font(bold=True)


def _add_availability_sheet(workbook: Workbook, messages: list[str], period_text: str) -> None:
    """Add explicit source limitations instead of silently filling missing data."""
    worksheet = workbook.create_sheet("Data Availability")
    worksheet.append(["Reporting Basis", period_text])
    worksheet.append(["Reporting Currency", "INR"])
    worksheet.append([])
    worksheet.append(["Data Availability Notes"])
    for message in messages:
        worksheet.append([message])
    worksheet.column_dimensions["A"].width = 105
    worksheet.column_dimensions["B"].width = 28
    worksheet["A1"].font = Font(bold=True)
    worksheet["A2"].font = Font(bold=True)
    worksheet["A4"].font = Font(bold=True, color="FFFFFF")
    worksheet["A4"].fill = PatternFill("solid", fgColor="1F2D4E")
    for row_number in range(5, worksheet.max_row + 1):
        worksheet.cell(row_number, 1).alignment = Alignment(wrap_text=True, vertical="top")


def _write_consolidated_pl_workbook(
    destination_path: Path,
    report_period: dict[str, Any],
    values: dict[str, float],
    messages: list[str],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Consolidated P&L"
    rows = [
        ("Revenue", None),
        ("Revenue - Available Data", values["Revenue"]),
        ("COGS", None),
        ("COGS - Available Data", values["COGS"]),
        ("Gross Profit", "=B6-B8"),
        ("Operating Expenses", None),
        ("Operating Expenses - Available Data", values["Operating Expenses"]),
        ("Operating Profit / EBITDA", "=B9-B11"),
        ("Other Income", None),
        ("Other Income - Available Data", values["Other Income"]),
        ("Other Expenses", None),
        ("Other Expenses - Available Data", values["Other Expenses"]),
        ("Net Profit", "=B12+B14-B16"),
    ]
    for row_number, (label, value) in enumerate(rows, start=5):
        worksheet.cell(row_number, 1, label)
        worksheet.cell(row_number, 2, value)
    _style_consolidated_statement(
        worksheet,
        "MIDOFFICE DATA | Consolidated Profit & Loss Statement",
        f"{report_period['header_title']} | India + US | Reporting Currency: INR",
    )
    _add_availability_sheet(workbook, messages, f"{report_period['start_date']} to {report_period['end_date']}")
    workbook.calculation.forceFullCalc = True
    workbook.calculation.fullCalcOnLoad = True
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination_path)


def _write_consolidated_balance_sheet_workbook(
    destination_path: Path,
    report_period: dict[str, Any],
    values: dict[str, float | None],
    messages: list[str],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Consolidated Balance Sheet"
    rows = [
        ("Assets", None),
        ("Current Assets", "=SUM(B7:B9)"),
        ("Bank / Cash", values["Bank / Cash"]),
        ("Receivables", values["Receivables"]),
        ("Other Assets", values["Other Assets"]),
        ("Total Assets", "=B6"),
        ("Liabilities", None),
        ("Payables", values["Payables"]),
        ("Tax Liabilities", values["Tax Liabilities"]),
        ("Total Liabilities", "=SUM(B12:B13)"),
        ("Equity", None),
        ("Equity / Retained Earnings", values["Equity / Retained Earnings"]),
        ("Total Liabilities and Equity", "=B14+B16"),
        ("Balance Check", "=B10-B17"),
    ]
    for row_number, (label, value) in enumerate(rows, start=5):
        worksheet.cell(row_number, 1, label)
        worksheet.cell(row_number, 2, value)
    _style_consolidated_statement(
        worksheet,
        "MIDOFFICE DATA | Consolidated Balance Sheet",
        f"As of {report_period['end_date'].strftime('%d %b %Y')} | India + US | Reporting Currency: INR",
    )
    _add_availability_sheet(workbook, messages, f"As of {report_period['end_date']}")
    workbook.calculation.forceFullCalc = True
    workbook.calculation.fullCalcOnLoad = True
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination_path)


def generate_consolidated_pl_report(
    financial_year: int | str,
    output_dir: str | Path,
    project_id: str | None = None,
    period_type: str = "full_year",
    selected_quarter: str | None = None,
    selected_half: str | None = None,
    monthly_pl_df: pd.DataFrame | None = None,
) -> dict[str, Any]:
    """Generate an INR consolidated P&L for India and US organizations."""
    financial_year_start = parse_financial_year_start(financial_year)
    report_period = get_consolidated_report_period(
        financial_year_start,
        period_type,
        selected_quarter=selected_quarter,
        selected_half=selected_half,
    )
    if monthly_pl_df is None:
        try:
            monthly_pl_df, _ = fetch_gold_mis_data(project_id, org_filter="all", report_period=report_period)
        except Exception:
            monthly_pl_df = pd.DataFrame()
    monthly_pl_df = _filter_dataframe_by_date_range(
        monthly_pl_df if monthly_pl_df is not None else pd.DataFrame(),
        "report_month",
        report_period["start_date"],
        report_period["end_date"],
    )
    values, messages = _consolidated_pl_values(monthly_pl_df)
    report_path = Path(output_dir) / (
        f"Consolidated_PL_{_financial_year_token(financial_year_start)}_{report_period['file_suffix']}.xlsx"
    )
    _write_consolidated_pl_workbook(report_path, report_period, values, messages)
    preview = pd.DataFrame([{"Line Item": label, "Amount (INR)": amount} for label, amount in values.items()])
    return {
        "status": "success",
        "message": f"Consolidated P&L generated for {report_period['period_name']}.",
        "report_type": "Consolidated P&L",
        "report_path": report_path,
        "report_period": report_period,
        "metrics": {
            "Revenue": _format_currency(values["Revenue"]),
            "COGS": _format_currency(values["COGS"]),
            "Operating Profit / EBITDA": _format_currency(values["Operating Profit / EBITDA"]),
            "Net Profit": _format_currency(values["Net Profit"]),
            "Reporting Currency": "INR",
        },
        "report_preview": preview,
        "data_message": " ".join(messages),
        "is_placeholder": False,
    }


def generate_consolidated_balance_sheet_report(
    financial_year: int | str,
    output_dir: str | Path,
    project_id: str | None = None,
    period_type: str = "full_year",
    selected_quarter: str | None = None,
    selected_half: str | None = None,
    balance_sheet_df: pd.DataFrame | None = None,
) -> dict[str, Any]:
    """Generate an INR consolidated balance sheet at the selected period end."""
    financial_year_start = parse_financial_year_start(financial_year)
    report_period = get_consolidated_report_period(
        financial_year_start,
        period_type,
        selected_quarter=selected_quarter,
        selected_half=selected_half,
    )
    if balance_sheet_df is None:
        try:
            balance_sheet_df = fetch_consolidated_balance_sheet_data(project_id, report_period["end_date"])
        except Exception:
            balance_sheet_df = pd.DataFrame()
    balance_sheet_df = _filter_balance_sheet_as_of(
        balance_sheet_df if balance_sheet_df is not None else pd.DataFrame(),
        report_period["end_date"],
    )
    values, messages = _balance_sheet_values(balance_sheet_df)
    report_path = Path(output_dir) / (
        f"Consolidated_Balance_Sheet_{_financial_year_token(financial_year_start)}_{report_period['file_suffix']}.xlsx"
    )
    _write_consolidated_balance_sheet_workbook(report_path, report_period, values, messages)
    assets = sum(float(values[label] or 0) for label in ["Bank / Cash", "Receivables", "Other Assets"])
    liabilities = sum(float(values[label] or 0) for label in ["Payables", "Tax Liabilities"])
    equity = float(values["Equity / Retained Earnings"] or 0)
    preview = pd.DataFrame([{"Line Item": label, "Amount (INR)": amount} for label, amount in values.items()])
    return {
        "status": "success",
        "message": f"Consolidated Balance Sheet generated as of {report_period['end_date']}.",
        "report_type": "Consolidated Balance Sheet",
        "report_path": report_path,
        "report_period": report_period,
        "as_of_date": report_period["end_date"],
        "metrics": {
            "Assets": _format_currency(assets),
            "Liabilities": _format_currency(liabilities),
            "Equity / Retained Earnings": _format_currency(equity),
            "As Of Date": str(report_period["end_date"]),
            "Reporting Currency": "INR",
        },
        "report_preview": preview,
        "data_message": " ".join(messages),
        "is_placeholder": False,
    }


def main() -> None:
    """CLI entrypoint for local MIS workbook generation."""
    project_root = _repo_root()
    result = generate_mis_report(2025, project_root / "outputs")
    print(f"Generated MIS report: {result['report_path']}")
