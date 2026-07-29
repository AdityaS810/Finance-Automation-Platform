"""Tests for the simplified Phase 5 Vendor Reconciliation workbook."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from google.auth.exceptions import RefreshError
from openpyxl import load_workbook

from backend.reports import vendor_transactions_report as vendor_report
from backend.reports.vendor_transactions_report import (
    ANALYSIS_DISCLAIMER,
    BANK_TRANSACTION_COLUMNS,
    DATASET_COLUMNS,
    INPUT_BANK_STATEMENT_COLUMNS,
    INPUT_DATASET_COLUMNS,
    INPUT_PAYMENT_ALLOCATION_COLUMNS,
    INPUT_VENDOR_BILL_COLUMNS,
    INPUT_VENDOR_PAYMENT_COLUMNS,
    INPUT_WORKBOOK_SHEETS,
    PAYMENT_ALLOCATION_COLUMNS,
    REPORT_MODE_ANALYSIS,
    REPORT_MODE_INPUT,
    RECONCILIATION_RESULT_COLUMNS,
    VENDOR_BILL_COLUMNS,
    VENDOR_PAYMENT_COLUMNS,
    WORKBOOK_SHEETS,
    ZOHO_PAYMENT_BATCH_COLUMNS,
    VendorReportValidationError,
    _query_parameters,
    _query_to_dataframe,
    build_reconciliation_input_queries,
    build_vendor_reconciliation_queries,
    combine_zoho_payment_batch,
    create_reconciliation_input_workbook,
    create_vendor_report_workbook,
    generate_vendor_transactions_report,
    safe_exception_details,
    summarize_reconciliation_input,
    summarize_vendor_report,
    validate_reconciliation_input_data,
    validate_vendor_report_data,
    user_facing_report_error,
)


def _sample_datasets() -> dict[str, pd.DataFrame]:
    datasets = {
        name: pd.DataFrame(columns=columns)
        for name, columns in DATASET_COLUMNS.items()
    }
    datasets["vendor_bills"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Bill ID": "bill-000000000001",
                "Bill Number": "B-100",
                "Bill Date": date(2026, 7, 1),
                "Due Date": date(2026, 7, 31),
                "Currency": "INR",
                "Taxable Amount": 90,
                "GST Amount": 10,
                "Bill Total": 100,
                "Outstanding Balance": 0,
                "Source Bill Status": "paid",
                "Allocated Payment Count": 1,
                "Allocated Amount": 100,
                "Bank-Verified Amount": 0,
                "Remaining Reconciliation Amount": 100,
                "Reconciliation Status": "Payment Recorded - Bank Pending",
                "Reconciliation Reason": "Payment is pending bank review",
                "Review Required": True,
                "Source Record ID": "source-bill-1",
                "Data Quality Status": "Valid",
            }
        ],
        columns=VENDOR_BILL_COLUMNS,
    )
    datasets["vendor_payments"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Payment ID": "payment-000000000001",
                "Payment Number": "P-100",
                "Payment Date": date(2026, 7, 2),
                "Reference Number": "REF-1",
                "Payment Mode": "banktransfer",
                "Currency": "INR",
                "Payment Amount": 100,
                "Allocated Amount": 100,
                "Unapplied Amount": 0,
                "Allocation Status": "Fully Allocated",
                "Paid-Through Account Name": "Operating Account",
                "Bank Match Status": "Bank Match Pending Review",
                "Bank Match Method": "exact_account_amount_date_window_review",
                "Bank Match Reason": "Manual review required",
                "Bank Transaction Date": date(2026, 7, 3),
                "Bank Amount": 100,
                "Review Required": True,
                "Source Record ID": "source-payment-1",
                "Bank Transaction Leg Key": "leg-debit",
                "Data Quality Status": "Fully Allocated",
            }
        ],
        columns=VENDOR_PAYMENT_COLUMNS,
    )
    datasets["payment_allocations"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Payment ID": "payment-000000000001",
                "Bill ID": "bill-000000000001",
                "Bill Number": "B-100",
                "Bill Payment ID": "bill-payment-1",
                "Currency": "INR",
                "Amount Applied": 100,
                "Payment Date": date(2026, 7, 2),
                "Allocation Key": "allocation-1",
                "Source Record ID": "source-payment-1",
                "Data Quality Status": "Valid",
            }
        ],
        columns=PAYMENT_ALLOCATION_COLUMNS,
    )
    datasets["bank_transactions"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Account Label": "Operating Account",
                "Transaction ID": "transaction-1",
                "Transaction Date": date(2026, 7, 3),
                "Transaction Type": "transfer",
                "Debit/Credit": "debit",
                "Direction": "Outgoing",
                "Currency": "INR",
                "Amount": 100,
                "Signed Amount": -100,
                "Status": "cleared",
                "Multi-Leg Transaction": True,
                "Used in Vendor Match": True,
                "Review Required": True,
                "Source Record ID": "source-bank-1",
                "Bank Transaction Leg Key": "leg-debit",
                "Match Method": "bank_transaction_leg",
                "Data Quality Status": "Valid",
            },
            {
                "Organization": "india",
                "Account Label": "Savings Account",
                "Transaction ID": "transaction-1",
                "Transaction Date": date(2026, 7, 3),
                "Transaction Type": "transfer",
                "Debit/Credit": "credit",
                "Direction": "Incoming",
                "Currency": "INR",
                "Amount": 100,
                "Signed Amount": 100,
                "Status": "cleared",
                "Multi-Leg Transaction": True,
                "Used in Vendor Match": False,
                "Review Required": False,
                "Source Record ID": "source-bank-1",
                "Bank Transaction Leg Key": "leg-credit",
                "Match Method": "bank_transaction_leg",
                "Data Quality Status": "Valid",
            },
        ],
        columns=BANK_TRANSACTION_COLUMNS,
    )
    datasets["reconciliation_results"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Bill ID": "bill-000000000001",
                "Bill Number": "B-100",
                "Bill Date": date(2026, 7, 1),
                "Due Date": date(2026, 7, 31),
                "Currency": "INR",
                "Bill Amount": 100,
                "Source Outstanding Balance": 0,
                "Source Bill Status": "paid",
                "Payment Count": 1,
                "Payment IDs": "payment-000000000001",
                "Latest Payment Date": date(2026, 7, 2),
                "Payment Amount": 100,
                "Allocated Amount": 100,
                "Bank Transaction Date": date(2026, 7, 3),
                "Bank Amount": 100,
                "Bank Match Status": "Bank Match Pending Review",
                "Bank Match Method": "exact_account_amount_date_window_review",
                "Bank Pending Amount": 100,
                "Remaining Reconciliation Amount": 100,
                "Reconciliation Status": "Payment Recorded - Bank Pending",
                "Reconciliation Reason": "Payment is pending bank review",
                "Exception Type": "Payment without bank match",
                "Review Required": True,
                "Review Status": "Pending",
                "Reviewer Comment": None,
            }
        ],
        columns=RECONCILIATION_RESULT_COLUMNS,
    )
    return datasets


def _sample_input_datasets() -> dict[str, pd.DataFrame]:
    datasets = {
        name: pd.DataFrame(columns=columns)
        for name, columns in INPUT_DATASET_COLUMNS.items()
    }
    datasets["vendor_bills"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Bill ID": "bill-000000000001",
                "Bill Number": "B-100",
                "Bill Date": date(2026, 7, 1),
                "Due Date": date(2026, 7, 31),
                "Currency": "INR",
                "Bill Amount": 100,
                "Outstanding Balance": 0,
                "Zoho Bill Status": "paid",
            },
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Bill ID": "bill-000000000002",
                "Bill Number": "B-101",
                "Bill Date": date(2026, 7, 2),
                "Due Date": date(2026, 8, 1),
                "Currency": "INR",
                "Taxable Amount": 45,
                "GST Amount": 5,
                "Bill Total": 50,
                "Outstanding Balance": 0,
                "Zoho Bill Status": "paid",
            },
        ],
        columns=INPUT_VENDOR_BILL_COLUMNS,
    )
    datasets["vendor_payments"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Vendor ID": "vendor-1",
                "Vendor Name": "Example Vendor",
                "Payment ID": "payment-000000000001",
                "Payment Number": "P-100",
                "Payment Date": date(2026, 7, 3),
                "Payment Reference": "REF-1",
                "Payment Mode": "banktransfer",
                "Paid-Through Account Name": "Operating Account",
                "Currency": "INR",
                "Total Payment Amount": 150,
                "Allocated Amount": 150,
                "Unapplied Amount": 0,
                "Allocation Status": "Fully Allocated",
            },
            {
                "Organization": "india",
                "Vendor ID": "vendor-2",
                "Vendor Name": "Unallocated Vendor",
                "Payment ID": "payment-000000000002",
                "Payment Number": "P-101",
                "Payment Date": date(2026, 7, 4),
                "Payment Reference": "REF-2",
                "Payment Mode": "banktransfer",
                "Paid-Through Account Name": "Operating Account",
                "Currency": "INR",
                "Total Payment Amount": 25,
                "Allocated Amount": 0,
                "Unapplied Amount": 25,
                "Allocation Status": "Unallocated",
            },
        ],
        columns=INPUT_VENDOR_PAYMENT_COLUMNS,
    )
    datasets["payment_allocations"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Payment ID": "payment-000000000001",
                "Bill Payment ID": "bill-payment-1",
                "Bill ID": "bill-000000000001",
                "Bill Number": "B-100",
                "Amount Applied": 100,
                "Allocation Key": "allocation-1",
            },
            {
                "Organization": "india",
                "Payment ID": "payment-000000000001",
                "Bill Payment ID": "bill-payment-2",
                "Bill ID": "bill-000000000002",
                "Bill Number": "B-101",
                "Amount Applied": 50,
                "Allocation Key": "allocation-2",
            },
        ],
        columns=INPUT_PAYMENT_ALLOCATION_COLUMNS,
    )
    datasets["bank_statement"] = pd.DataFrame(
        [
            {
                "Organization": "india",
                "Bank Source": "Uploaded Bank",
                "Upload ID": "upload-full-period",
                "Source File": "statement.xlsx",
                "Transaction Date": date(2026, 7, 3),
                "Value Date": date(2026, 7, 3),
                "Narration": "Vendor transfer for invoice B-100",
                "Reference Number": "REF-1",
                "Debit Amount": 100,
                "Credit Amount": 0,
                "Signed Amount": -100,
                "Direction": "Outgoing",
                "Currency": "INR",
                "Account Name": "Uploaded Bank",
                "Masked Account Number": "******1234",
                "Statement Row ID": "statement-row-debit",
            },
            {
                "Organization": "india",
                "Bank Source": "Uploaded Bank",
                "Upload ID": "upload-full-period",
                "Source File": "statement.xlsx",
                "Transaction Date": date(2026, 7, 3),
                "Value Date": date(2026, 7, 3),
                "Narration": "Incoming transfer",
                "Reference Number": "REF-1",
                "Debit Amount": 0,
                "Credit Amount": 100,
                "Signed Amount": 100,
                "Direction": "Incoming",
                "Currency": "INR",
                "Account Name": "Uploaded Bank",
                "Masked Account Number": "******1234",
                "Statement Row ID": "statement-row-credit",
            },
        ],
        columns=INPUT_BANK_STATEMENT_COLUMNS,
    )
    return datasets


def _sample_input_datasets_with_older_bill() -> dict[str, pd.DataFrame]:
    datasets = _sample_input_datasets()
    older_bill = datasets["vendor_bills"].iloc[1].copy()
    older_bill["Bill ID"] = "bill-older-0000000001"
    older_bill["Bill Number"] = "Nov 2025"
    older_bill["Bill Date"] = date(2025, 11, 30)
    older_bill["Due Date"] = date(2025, 12, 30)
    datasets["current_bills"] = pd.concat(
        [
            datasets["vendor_bills"].reindex(
                columns=vendor_report.INPUT_VENDOR_BILL_SOURCE_COLUMNS
            ),
            pd.DataFrame([older_bill]).reindex(
                columns=vendor_report.INPUT_VENDOR_BILL_SOURCE_COLUMNS
            ),
        ],
        ignore_index=True,
    )
    datasets["payment_allocations"].loc[1, "Bill ID"] = older_bill["Bill ID"]
    datasets["payment_allocations"].loc[1, "Bill Number"] = older_bill["Bill Number"]
    return datasets


def test_required_workbook_sheets_and_columns():
    workbook = create_vendor_report_workbook(_sample_datasets())

    assert len(WORKBOOK_SHEETS) == 6
    assert workbook.sheetnames == WORKBOOK_SHEETS
    assert "Vendor Master" not in workbook.sheetnames
    assert "Exceptions - Review" not in workbook.sheetnames
    assert "Technical Audit" not in workbook.sheetnames
    assert [cell.value for cell in workbook["Vendor Bills"][3]] == VENDOR_BILL_COLUMNS
    assert [cell.value for cell in workbook["Vendor Payments"][3]] == VENDOR_PAYMENT_COLUMNS
    assert [cell.value for cell in workbook["Reconciliation Results"][3]] == (
        RECONCILIATION_RESULT_COLUMNS
    )


def test_report_mode_selector_defaults_to_input_and_analysis_remains_available():
    page_source = (
        Path(__file__).parents[1] / "pages" / "06_Downloads.py"
    ).read_text(encoding="utf-8")

    assert 'st.selectbox(\n    "Report mode"' in page_source
    assert "options=REPORT_MODES" in page_source
    assert "index=0" in page_source
    assert REPORT_MODE_INPUT in page_source
    assert REPORT_MODE_ANALYSIS in page_source


def test_input_workbook_has_exactly_four_sheets_and_no_analysis_statuses():
    workbook = create_reconciliation_input_workbook(
        _sample_input_datasets(),
        filters={"start_date": date(2026, 7, 1), "end_date": date(2026, 7, 31)},
    )

    assert workbook.sheetnames == INPUT_WORKBOOK_SHEETS
    assert [cell.value for cell in workbook["Vendor Bills"][3]] == (
        INPUT_VENDOR_BILL_COLUMNS
    )
    assert [cell.value for cell in workbook["Zoho Payment Batch"][3]] == (
        ZOHO_PAYMENT_BATCH_COLUMNS
    )
    workbook_text = " ".join(
        str(cell.value or "")
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    assert "Reconciliation Status" not in workbook_text
    assert "Bank Match Status" not in workbook_text
    assert "Review Required" not in workbook_text


def test_payment_batch_combines_allocations_and_keeps_unallocated_payment():
    datasets = _sample_input_datasets()
    batch = combine_zoho_payment_batch(
        datasets["vendor_payments"],
        datasets["payment_allocations"],
    )

    allocated_rows = batch[
        batch["Payment ID"].eq("payment-000000000001")
    ]
    unallocated_rows = batch[
        batch["Payment ID"].eq("payment-000000000002")
    ]
    assert len(allocated_rows) == 2
    assert set(allocated_rows["Bill ID"]) == {
        "bill-000000000001",
        "bill-000000000002",
    }
    assert len(unallocated_rows) == 1
    assert pd.isna(unallocated_rows.iloc[0]["Bill ID"])


def test_input_statement_rows_remain_distinct_and_ids_are_excel_text():
    workbook = create_reconciliation_input_workbook(_sample_input_datasets())
    bank_sheet = workbook["Bank Statement"]
    payment_sheet = workbook["Zoho Payment Batch"]

    assert bank_sheet.max_row == 5
    assert {bank_sheet["P4"].value, bank_sheet["P5"].value} == {
        "statement-row-debit",
        "statement-row-credit",
    }
    assert bank_sheet["C4"].number_format == "@"
    assert bank_sheet["P4"].number_format == "@"
    assert payment_sheet["D4"].number_format == "@"
    assert payment_sheet["H4"].number_format == "@"


def test_manual_input_sheet_columns_match_finance_review_contract():
    assert INPUT_VENDOR_BILL_COLUMNS[:13] == [
        "Organization",
        "Vendor ID",
        "Vendor Name",
        "Bill ID",
        "Bill Number",
        "Bill Date",
        "Due Date",
        "Currency",
        "Taxable Amount",
        "GST Amount",
        "Bill Total",
        "Outstanding Balance",
        "Zoho Bill Status",
    ]
    assert ZOHO_PAYMENT_BATCH_COLUMNS == [
        "Organization",
        "Vendor ID",
        "Vendor Name",
        "Payment ID",
        "Payment Date",
        "Payment Reference",
        "Payment Amount",
        "Bill ID",
        "Bill Number",
        "Amount Applied",
        "Unapplied Amount",
    ]
    assert INPUT_BANK_STATEMENT_COLUMNS[-7:] == [
        "Extracted Invoice Number",
        "Invoice Reference Result",
        "Matching Bill Number",
        "Matching Bill ID",
        "Possible Amount Pattern",
        "Amount Difference",
        "Manual Review Note",
    ]


def test_invoice_reference_tds_and_split_suggestions_are_conservative():
    template_bill = _sample_input_datasets()["vendor_bills"].iloc[0].copy()
    bills = []
    for number in ("INV-EXACT", "INV-TDS2", "INV-TDS10", "INV-SPLIT"):
        bill = template_bill.copy()
        bill["Bill ID"] = f"id-{number}"
        bill["Bill Number"] = number
        bill["Taxable Amount"] = 100
        bill["GST Amount"] = 18
        bill["Bill Total"] = 118
        bills.append(bill)
    bill_frame = pd.DataFrame(bills, columns=INPUT_VENDOR_BILL_COLUMNS)

    template_bank = _sample_input_datasets()["bank_statement"].iloc[0].copy()
    bank_specs = [
        ("exact", "Paid INV-EXACT in full", "", 118),
        ("tds2", "Paid against invoice inv/tds2", "", 116),
        ("tds10", "Vendor payment", "INV_TDS10", 108),
        ("split-a", "Part one INV-SPLIT", "", 50),
        ("split-b", "Part two INV-SPLIT", "", 68),
        ("none", "General vendor settlement 2026-03", "BANK-REF", 75),
    ]
    bank_rows = []
    original_descriptions = []
    for key, description, reference, amount in bank_specs:
        row = template_bank.copy()
        row["Statement Row ID"] = f"row-{key}"
        row["Narration"] = description
        row["Reference Number"] = reference
        row["Debit Amount"] = amount
        row["Credit Amount"] = 0
        row["Signed Amount"] = -amount
        bank_rows.append(row)
        original_descriptions.append(description)
    bank_frame = pd.DataFrame(
        bank_rows,
        columns=vendor_report.INPUT_BANK_STATEMENT_SOURCE_COLUMNS,
    )

    reviewed = vendor_report.prepare_manual_review_bank_statement(
        bill_frame,
        bank_frame,
    )
    by_key = reviewed.set_index("Statement Row ID")

    assert reviewed["Narration"].tolist() == original_descriptions
    assert by_key.loc["row-exact", "Possible Amount Pattern"] == (
        "Exact Invoice Total Candidate"
    )
    assert by_key.loc["row-tds2", "Possible Amount Pattern"] == (
        "Possible 2% TDS Candidate"
    )
    assert by_key.loc["row-tds10", "Possible Amount Pattern"] == (
        "Possible 10% TDS Candidate"
    )
    assert by_key.loc["row-split-a", "Possible Amount Pattern"] == (
        "Possible Split Payment - Review Required"
    )
    assert by_key.loc["row-split-b", "Possible Amount Pattern"] == (
        "Possible Split Payment - Review Required"
    )
    assert by_key.loc["row-none", "Invoice Reference Result"] == (
        "No Invoice Reference"
    )
    assert pd.isna(by_key.loc["row-none", "Extracted Invoice Number"])

    summary = summarize_reconciliation_input(
        {
            "vendor_bills": bill_frame,
            "vendor_payments": pd.DataFrame(columns=INPUT_VENDOR_PAYMENT_COLUMNS),
            "payment_allocations": pd.DataFrame(
                columns=INPUT_PAYMENT_ALLOCATION_COLUMNS
            ),
            "bank_statement": reviewed,
        }
    )
    assert summary["bank_debit_count"] == 6
    assert summary["descriptions_with_invoice_references"] == 5
    assert summary["exact_invoice_total_candidates"] == 1
    assert summary["possible_2_percent_tds_candidates"] == 1
    assert summary["possible_10_percent_tds_candidates"] == 1
    assert summary["possible_split_payment_cases"] == 1
    assert summary["records_with_no_invoice_reference"] == 1


def test_upload_recommendation_uses_coverage_not_latest_and_never_merges():
    uploads = [
        {
            "upload_id": "latest-partial",
            "coverage_start_date": date(2026, 3, 1),
            "coverage_end_date": date(2026, 3, 31),
            "period_row_count": 10,
        },
        {
            "upload_id": "older-full-period",
            "coverage_start_date": date(2025, 12, 1),
            "coverage_end_date": date(2026, 4, 1),
            "period_row_count": 100,
        },
    ]

    assert vendor_report.recommend_bank_statement_upload(
        uploads,
        date(2026, 1, 1),
        date(2026, 3, 31),
    ) == "older-full-period"
    assert vendor_report.recommend_bank_statement_upload(
        uploads[:1],
        date(2026, 1, 1),
        date(2026, 3, 31),
    ) is None


def test_downloads_page_has_explicit_upload_selector_and_no_merge_message():
    page_source = (
        Path(__file__).parents[1] / "pages" / "06_Downloads.py"
    ).read_text(encoding="utf-8")

    assert '"Bank statement upload"' in page_source
    assert "recommend_bank_statement_upload(" in page_source
    assert "uploads will not be merged" in page_source
    assert "bank_upload_id=selected_bank_upload_id" in page_source


def test_reference_classifications_reject_unsafe_candidates_and_only_match_unique():
    datasets = _sample_input_datasets()
    bill_template = datasets["vendor_bills"].iloc[0].copy()
    bills = []
    for bill_id, bill_number in (
        ("bill-unique", "INV-77881"),
        ("bill-ambiguous-a", "PO-99110"),
        ("bill-ambiguous-b", "AB-88220"),
        ("bill-short", "17"),
        ("bill-month", "Nov 2025"),
    ):
        row = bill_template.copy()
        row["Bill ID"] = bill_id
        row["Bill Number"] = bill_number
        row["Taxable Amount"] = 100
        row["Bill Total"] = 118
        bills.append(row)
    bill_frame = pd.DataFrame(bills, columns=INPUT_VENDOR_BILL_COLUMNS)

    statement_template = datasets["bank_statement"].iloc[0].copy()
    statement_specs = (
        ("unique", "PAYMENT INV / 77881 CONFIRMED", "", 118),
        ("embedded", "NEFTXXINV77881SETTLEMENT", "", 117.78),
        ("secondary", "VENDOR PAYMENT", "INV-77881", 116),
        ("ambiguous", "PAY PO-99110 AND AB/88220", "", 118),
        ("short", "PAYMENT FOR BILL 17", "", 118),
        ("month", "NOV 2025 SERVICES", "", 118),
        ("year", "GENERAL PAYMENT 2026", "", 118),
        ("none", "GENERAL VENDOR PAYMENT", "", 118),
    )
    statement_rows = []
    for key, narration, reference, debit in statement_specs:
        row = statement_template.copy()
        row["Statement Row ID"] = key
        row["Narration"] = narration
        row["Reference Number"] = reference
        row["Debit Amount"] = debit
        row["Signed Amount"] = -debit
        statement_rows.append(row)
    reviewed = vendor_report.prepare_manual_review_bank_statement(
        bill_frame,
        pd.DataFrame(
            statement_rows,
            columns=vendor_report.INPUT_BANK_STATEMENT_SOURCE_COLUMNS,
        ),
    ).set_index("Statement Row ID")

    assert reviewed.loc["unique", "Invoice Reference Result"] == (
        "Unique High-Confidence Reference"
    )
    assert reviewed.loc["embedded", "Invoice Reference Result"] == (
        "Unique High-Confidence Reference"
    )
    assert reviewed.loc["secondary", "Invoice Reference Result"] == (
        "Unique High-Confidence Reference"
    )
    assert reviewed.loc["ambiguous", "Invoice Reference Result"] == (
        "Ambiguous Reference"
    )
    assert reviewed.loc["short", "Invoice Reference Result"] == (
        "Low-Specificity Candidate"
    )
    assert reviewed.loc["month", "Invoice Reference Result"] == (
        "Low-Specificity Candidate"
    )
    assert reviewed.loc["year", "Invoice Reference Result"] == (
        "No Invoice Reference"
    )
    assert reviewed.loc["none", "Invoice Reference Result"] == (
        "No Invoice Reference"
    )
    assert reviewed.loc["unique", "Matching Bill ID"] == "bill-unique"
    assert reviewed.loc[
        ["ambiguous", "short", "month", "year", "none"],
        "Matching Bill ID",
    ].isna().all()
    assert reviewed.loc["secondary", "Possible Amount Pattern"] == (
        "Possible 2% TDS Candidate"
    )
    assert reviewed.loc["short", "Possible Amount Pattern"] == (
        "No Unique Invoice Reference"
    )


def test_amount_tolerance_does_not_call_point_two_two_variance_exact():
    datasets = _sample_input_datasets()
    bill = datasets["vendor_bills"].iloc[[0]].copy()
    bill.loc[:, "Bill Number"] = "INV-44001"
    bill.loc[:, "Bill Total"] = 100
    bill.loc[:, "Taxable Amount"] = 80
    statement = datasets["bank_statement"].iloc[[0]].copy()
    statement.loc[:, "Narration"] = "Paid INV-44001"
    statement.loc[:, "Debit Amount"] = 100.22
    statement.loc[:, "Signed Amount"] = -100.22

    reviewed = vendor_report.prepare_manual_review_bank_statement(bill, statement)

    assert reviewed.loc[0, "Possible Amount Pattern"] == (
        "Amount Difference - Review Required"
    )
    assert reviewed.loc[0, "Amount Difference"] == pytest.approx(0.22)


@pytest.mark.parametrize(
    ("mutation", "expected_failure"),
    [
        ("duplicate-row", "duplicate_statement_row_keys=2"),
        ("mixed-upload", "selected_bank_upload_count=2"),
        ("unmasked-account", "unmasked_account_numbers=1"),
    ],
)
def test_statement_safety_validations_block_unsafe_exports(
    mutation,
    expected_failure,
):
    datasets = _sample_input_datasets()
    if mutation == "duplicate-row":
        datasets["bank_statement"].loc[1, "Statement Row ID"] = (
            datasets["bank_statement"].loc[0, "Statement Row ID"]
        )
    elif mutation == "mixed-upload":
        datasets["bank_statement"].loc[1, "Upload ID"] = "another-upload"
    else:
        datasets["bank_statement"].loc[0, "Masked Account Number"] = "1234567890"

    with pytest.raises(VendorReportValidationError, match=expected_failure):
        create_reconciliation_input_workbook(datasets)


def test_input_queries_are_silver_only_and_parameterized():
    filters = {
        "start_date": date(2026, 7, 1),
        "end_date": date(2026, 7, 31),
        "organization": "india",
        "vendor": "Example Vendor",
        "currency": "INR",
        "bank_upload_id": "upload-full-period",
    }
    queries = build_reconciliation_input_queries("project-1", filters)
    query_text = "\n".join(queries.values())

    assert set(queries) == set(INPUT_DATASET_COLUMNS) | {"current_bills"}
    assert "finance_silver.fact_bills" in query_text
    assert "finance_silver.fact_vendor_payments" in query_text
    assert "finance_silver.bridge_vendor_payment_bill_allocations" in query_text
    assert "finance_silver.fact_bank_statement_lines" in query_text
    assert "finance_silver.fact_bank_transactions" not in query_text
    assert "finance_silver.fact_transactions" not in query_text
    assert "finance_gold." not in query_text
    assert "@start_date" in query_text
    assert "@organization" in query_text
    assert "@bank_upload_id" in query_text
    assert "'NOT PROVIDED'" in queries["bank_statement"]
    assert "Example Vendor" not in query_text
    assert "SELECT DISTINCT" in queries["current_bills"]
    assert "INNER JOIN `project-1.finance_silver.bridge_vendor_payment_bill_allocations`" in (
        queries["current_bills"]
    )
    assert "@start_date" in queries["current_bills"]
    assert "@end_date" in queries["current_bills"]
    assert "bill.taxable_amount" in queries["vendor_bills"]
    assert "bill.tax_amount AS gst_amount" in queries["vendor_bills"]


def test_input_bill_query_mapping_preserves_tax_and_total_amounts():
    raw = pd.DataFrame(
        [
            {
                "organization": "india",
                "vendor_id": "vendor-1",
                "vendor_name": "Example Vendor",
                "bill_id": "bill-1",
                "bill_number": "INV-1",
                "bill_date": date(2026, 3, 1),
                "due_date": date(2026, 3, 31),
                "currency": "INR",
                "taxable_amount": 100,
                "gst_amount": 18,
                "bill_total": 118,
                "outstanding_balance": 0,
                "zoho_bill_status": "paid",
            }
        ]
    )

    mapped = vendor_report._rename_input_query_columns(raw, "vendor_bills")

    assert mapped.loc[0, "Taxable Amount"] == 100
    assert mapped.loc[0, "GST Amount"] == 18
    assert mapped.loc[0, "Bill Total"] == 118


def test_payment_inside_period_includes_older_bill_once_with_warning():
    datasets = _sample_input_datasets_with_older_bill()
    selected_bills = vendor_report._select_reconciliation_input_bills(datasets)
    validation_source = dict(datasets)
    validation_source["vendor_bills"] = selected_bills

    validations = validate_reconciliation_input_data(validation_source)
    older_rows = selected_bills[
        selected_bills["Bill ID"].eq("bill-older-0000000001")
    ]

    assert len(older_rows) == 1
    assert older_rows.iloc[0]["In Selected Bill Period"] == "No"
    assert older_rows.iloc[0]["Inclusion Reason"] == (
        "Linked to Payment in Selected Period"
    )
    assert validations["older_linked_bills_included"] == 1
    assert validations["informational_warnings"] == [
        "1 older bill was included because payments allocated to it fall "
        "within the selected report period."
    ]


def test_input_export_has_no_duplicate_bills_and_all_allocation_bills_exist():
    datasets = _sample_input_datasets_with_older_bill()
    workbook = create_reconciliation_input_workbook(
        datasets,
        filters={"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 27)},
    )
    bill_sheet = workbook["Vendor Bills"]
    payment_sheet = workbook["Zoho Payment Batch"]
    bill_headers = [cell.value for cell in bill_sheet[3]]
    payment_headers = [cell.value for cell in payment_sheet[3]]
    bill_id_column = bill_headers.index("Bill ID") + 1
    allocation_bill_id_column = payment_headers.index("Bill ID") + 1
    exported_bill_ids = {
        bill_sheet.cell(row, bill_id_column).value
        for row in range(4, bill_sheet.max_row + 1)
    }
    allocation_bill_ids = {
        payment_sheet.cell(row, allocation_bill_id_column).value
        for row in range(4, payment_sheet.max_row + 1)
        if payment_sheet.cell(row, allocation_bill_id_column).value
    }

    assert len(exported_bill_ids) == bill_sheet.max_row - 3
    assert allocation_bill_ids <= exported_bill_ids
    assert workbook.sheetnames == INPUT_WORKBOOK_SHEETS
    assert len(workbook.sheetnames) == 4
    summary_values = {
        workbook["Summary"].cell(row, 1).value: workbook["Summary"].cell(row, 2).value
        for row in range(1, workbook["Summary"].max_row + 1)
    }
    assert summary_values["Bills dated within selected period"] == 2
    assert summary_values["Older linked bills included"] == 1
    assert summary_values["Total bill rows exported"] == 3
    summary_labels = {
        workbook["Summary"].cell(row, 1).value: workbook["Summary"].cell(row, 2).value
        for row in range(3, workbook["Summary"].max_row + 1)
    }
    assert summary_labels["Information"] == (
        "1 older bill was included because payments allocated to it fall "
        "within the selected report period."
    )


def test_dynamic_warning_count_uses_generic_older_bill_fixtures():
    datasets = _sample_input_datasets()
    period_bills = datasets["vendor_bills"].reindex(
        columns=vendor_report.INPUT_VENDOR_BILL_SOURCE_COLUMNS
    )
    older_bills = []
    allocations = []
    for index in range(4):
        older_bill = period_bills.iloc[0].copy()
        older_bill["Bill ID"] = f"older-bill-{index}"
        older_bill["Bill Number"] = f"OLD-{index}"
        older_bill["Bill Date"] = date(2025, index + 1, 1)
        older_bills.append(older_bill)
        allocation = datasets["payment_allocations"].iloc[0].copy()
        allocation["Bill ID"] = older_bill["Bill ID"]
        allocation["Bill Number"] = older_bill["Bill Number"]
        allocation["Bill Payment ID"] = f"older-payment-{index}"
        allocation["Allocation Key"] = f"older-allocation-{index}"
        allocation["Amount Applied"] = 25
        allocations.append(allocation)
    datasets["current_bills"] = pd.concat(
        [period_bills, pd.DataFrame(older_bills)],
        ignore_index=True,
    )
    datasets["payment_allocations"] = pd.DataFrame(
        allocations,
        columns=INPUT_PAYMENT_ALLOCATION_COLUMNS,
    )
    datasets["vendor_payments"].loc[
        datasets["vendor_payments"]["Payment ID"].eq("payment-000000000001"),
        ["Total Payment Amount", "Allocated Amount"],
    ] = [100, 100]

    selected_bills = vendor_report._select_reconciliation_input_bills(datasets)
    validation_source = dict(datasets)
    validation_source["vendor_bills"] = selected_bills
    validations = validate_reconciliation_input_data(validation_source)

    assert selected_bills["In Selected Bill Period"].eq("No").sum() == 4
    assert validations["informational_warnings"] == [
        "4 older bills were included because payments allocated to them fall "
        "within the selected report period."
    ]


def test_blank_allocation_bill_id_is_preserved_and_does_not_block():
    datasets = _sample_input_datasets()
    blank_allocation = datasets["payment_allocations"].iloc[0].copy()
    blank_allocation["Bill ID"] = None
    blank_allocation["Bill Number"] = None
    blank_allocation["Bill Payment ID"] = "blank-bill-payment"
    blank_allocation["Allocation Key"] = "blank-bill-allocation"
    blank_allocation["Amount Applied"] = 0
    datasets["payment_allocations"] = pd.concat(
        [datasets["payment_allocations"], pd.DataFrame([blank_allocation])],
        ignore_index=True,
    )

    workbook = create_reconciliation_input_workbook(datasets)
    payment_sheet = workbook["Zoho Payment Batch"]
    headers = [cell.value for cell in payment_sheet[3]]
    payment_id_column = headers.index("Payment ID") + 1
    bill_id_column = headers.index("Bill ID") + 1
    blank_rows = [
        row
        for row in range(4, payment_sheet.max_row + 1)
        if payment_sheet.cell(row, payment_id_column).value
        == "payment-000000000001"
        and payment_sheet.cell(row, bill_id_column).value is None
    ]

    assert len(blank_rows) == 1
    assert payment_sheet.cell(blank_rows[0], bill_id_column).value is None


def test_empty_input_datasets_block_misleading_workbook_generation():
    datasets = {
        name: pd.DataFrame(columns=columns)
        for name, columns in INPUT_DATASET_COLUMNS.items()
    }

    with pytest.raises(
        VendorReportValidationError,
        match="statement_rows_present=False",
    ):
        create_reconciliation_input_workbook(datasets)


def test_genuinely_missing_allocation_bill_still_blocks_input_export():
    datasets = _sample_input_datasets_with_older_bill()
    datasets["payment_allocations"].loc[1, "Bill ID"] = "missing-bill-id"

    with pytest.raises(
        VendorReportValidationError,
        match="allocation_bill_ids_missing_from_bills=1",
    ):
        create_reconciliation_input_workbook(datasets)


def test_input_blocking_validations_and_empty_us_sources():
    datasets = _sample_input_datasets()
    usd_bank = datasets["bank_statement"].iloc[0].copy()
    usd_bank["Organization"] = "us"
    usd_bank["Statement Row ID"] = "us-statement-row-1"
    usd_bank["Currency"] = "USD"
    datasets["bank_statement"] = pd.concat(
        [datasets["bank_statement"], pd.DataFrame([usd_bank])],
        ignore_index=True,
    )

    validations = validate_reconciliation_input_data(datasets)
    summary = summarize_reconciliation_input(datasets)
    assert validations["organizations_kept_separate"] is True
    assert {row["currency"] for row in summary["currency_totals"]} == {
        "INR",
        "USD",
    }

    duplicate = datasets["payment_allocations"].iloc[[0]].copy()
    datasets["payment_allocations"] = pd.concat(
        [datasets["payment_allocations"], duplicate],
        ignore_index=True,
    )
    with pytest.raises(VendorReportValidationError, match="duplicate_allocation_keys"):
        create_reconciliation_input_workbook(datasets)


def test_analysis_summary_contains_finance_review_disclaimer():
    workbook = create_vendor_report_workbook(_sample_datasets())

    assert workbook["Summary"]["A2"].value == ANALYSIS_DISCLAIMER


def test_exception_review_fields_are_in_reconciliation_results():
    workbook = create_vendor_report_workbook(_sample_datasets())
    sheet = workbook["Reconciliation Results"]
    headers = [cell.value for cell in sheet[3]]
    record = dict(zip(headers, [cell.value for cell in sheet[4]]))

    assert headers[-4:] == [
        "Exception Type",
        "Review Required",
        "Review Status",
        "Reviewer Comment",
    ]
    assert record["Exception Type"] == "Payment without bank match"
    assert record["Review Status"] == "Pending"
    assert record["Reviewer Comment"] is None
    assert sheet.data_validations.dataValidation[0].formula1 == '"Pending,In Review,Resolved"'


def test_audit_fields_end_the_relevant_business_sheets():
    workbook = create_vendor_report_workbook(_sample_datasets())

    assert [cell.value for cell in workbook["Vendor Bills"][3]][-2:] == [
        "Source Record ID",
        "Data Quality Status",
    ]
    assert [cell.value for cell in workbook["Vendor Payments"][3]][-3:] == [
        "Source Record ID",
        "Bank Transaction Leg Key",
        "Data Quality Status",
    ]
    assert [cell.value for cell in workbook["Payment Allocations"][3]][-3:] == [
        "Allocation Key",
        "Source Record ID",
        "Data Quality Status",
    ]
    assert [cell.value for cell in workbook["Bank Transactions"][3]][-4:] == [
        "Source Record ID",
        "Bank Transaction Leg Key",
        "Match Method",
        "Data Quality Status",
    ]


def test_payment_allocations_do_not_duplicate_and_both_bank_legs_remain():
    datasets = _sample_datasets()
    validations = validate_vendor_report_data(datasets)
    workbook = create_vendor_report_workbook(datasets, validation_results=validations)

    assert validations["duplicate_allocation_keys"] == 0
    assert validations["duplicate_bank_transaction_leg_keys"] == 0
    assert workbook["Payment Allocations"].max_row == 4
    assert workbook["Bank Transactions"].max_row == 5
    assert {
        workbook["Bank Transactions"]["F4"].value,
        workbook["Bank Transactions"]["F5"].value,
    } == {"debit", "credit"}


def test_pending_review_records_remain_pending_and_no_reference_fallback():
    datasets = _sample_datasets()
    validations = validate_vendor_report_data(datasets)

    assert validations["pending_review_status_violations"] == 0
    assert datasets["vendor_payments"].iloc[0]["Bank Match Status"] == (
        "Bank Match Pending Review"
    )


def test_currency_totals_are_separate():
    datasets = _sample_datasets()
    usd = datasets["reconciliation_results"].iloc[0].copy()
    usd["Organization"] = "us"
    usd["Bill ID"] = "us-bill-1"
    usd["Currency"] = "USD"
    usd["Bill Amount"] = 25
    usd["Payment Amount"] = 25
    usd["Allocated Amount"] = 25
    usd["Bank Pending Amount"] = 25
    datasets["reconciliation_results"] = pd.concat(
        [datasets["reconciliation_results"], pd.DataFrame([usd])],
        ignore_index=True,
    )

    summary = summarize_vendor_report(datasets)

    assert {row["currency"] for row in summary["currency_totals"]} == {"INR", "USD"}
    assert len(summary["currency_totals"]) == 2


def test_empty_datasets_do_not_crash():
    datasets = {
        name: pd.DataFrame(columns=columns)
        for name, columns in DATASET_COLUMNS.items()
    }
    workbook = create_vendor_report_workbook(datasets)

    assert workbook.sheetnames == WORKBOOK_SHEETS
    assert [cell.value for cell in workbook["Vendor Payments"][3]] == VENDOR_PAYMENT_COLUMNS
    assert workbook["Summary"]["B9"].value == 0


def test_queries_use_verified_sources_and_parameterized_filters():
    filters = {
        "start_date": date(2026, 1, 1),
        "end_date": date(2026, 7, 31),
        "organization": "india",
        "vendor": "Example Vendor",
        "currency": "INR",
        "source_bill_status": "paid",
        "reconciliation_status": "Payment Recorded - Bank Pending",
        "bank_match_status": "Bank Match Pending Review",
        "review_required": True,
    }
    queries = build_vendor_reconciliation_queries("project-1", filters)
    query_text = "\n".join(queries.values())
    parameters = _query_parameters(**filters)

    assert set(queries) == set(DATASET_COLUMNS)
    assert "finance_silver.fact_transactions" not in query_text
    assert "finance_silver.fact_bank_transactions" in query_text
    assert "finance_gold.vendor_reconciliation_exceptions" in query_text
    assert "@start_date" in query_text
    assert "@organization" in query_text
    assert "@bank_match_status" in query_text
    assert "Example Vendor" not in query_text
    assert {parameter.name for parameter in parameters} == set(filters)


def test_vendor_filter_options_are_sourced_from_dim_contacts(monkeypatch):
    captured = {}

    def fake_query_to_dataframe(client, query, **kwargs):
        captured["query"] = query
        return pd.DataFrame(
            [
                {"option_type": "vendor", "option_value": "Misc"},
                {
                    "option_type": "vendor",
                    "option_value": "Midoffice Solutions Pvt Ltd",
                },
            ]
        )

    monkeypatch.setattr(vendor_report, "_query_to_dataframe", fake_query_to_dataframe)
    options = vendor_report.fetch_vendor_filter_options(
        project_id="project-1",
        client=object(),
    )

    assert options["vendors"] == ["Midoffice Solutions Pvt Ltd", "Misc"]
    assert "finance_silver.dim_contacts" in captured["query"]
    assert "LOWER(contact_type) = 'vendor'" in captured["query"]


def test_blocking_validation_stops_export():
    datasets = _sample_datasets()
    duplicate = datasets["payment_allocations"].iloc[[0]].copy()
    datasets["payment_allocations"] = pd.concat(
        [datasets["payment_allocations"], duplicate],
        ignore_index=True,
    )

    with pytest.raises(VendorReportValidationError, match="duplicate_allocation_keys"):
        create_vendor_report_workbook(datasets)


def test_allocation_exceeding_payment_stops_export():
    datasets = _sample_datasets()
    datasets["payment_allocations"].loc[0, "Amount Applied"] = 101

    with pytest.raises(VendorReportValidationError, match="allocations_exceeding_payment"):
        validate_vendor_report_data(datasets)


def test_credit_leg_cannot_be_used_in_vendor_match():
    datasets = _sample_datasets()
    datasets["bank_transactions"].loc[1, "Used in Vendor Match"] = True

    with pytest.raises(
        VendorReportValidationError,
        match="credit_legs_used_as_vendor_payments",
    ):
        validate_vendor_report_data(datasets)


def test_excel_identifiers_and_business_audit_fields_are_text():
    workbook = create_vendor_report_workbook(_sample_datasets())
    bills = workbook["Vendor Bills"]
    payments = workbook["Vendor Payments"]

    assert bills["B4"].number_format == "@"
    assert bills["D4"].number_format == "@"
    assert bills.cell(4, VENDOR_BILL_COLUMNS.index("Source Record ID") + 1).value == (
        "source-bill-1"
    )
    assert payments["D4"].number_format == "@"
    assert payments.cell(
        4,
        VENDOR_PAYMENT_COLUMNS.index("Bank Transaction Leg Key") + 1,
    ).number_format == "@"


def test_query_to_dataframe_preserves_empty_result_columns():
    class EmptyResult:
        schema = [
            SimpleNamespace(name="organization"),
            SimpleNamespace(name="vendor_name"),
        ]

        def __iter__(self):
            return iter(())

    class EmptyQueryJob:
        def result(self, timeout=None):
            return EmptyResult()

    class FakeClient:
        def query(self, query, **kwargs):
            return EmptyQueryJob()

    result = _query_to_dataframe(FakeClient(), "SELECT organization, vendor_name")

    assert result.empty
    assert list(result.columns) == ["organization", "vendor_name"]


def test_query_retries_once_with_bounded_timeouts():
    class EmptyResult:
        schema = [SimpleNamespace(name="organization")]

        def __iter__(self):
            return iter(())

    class TimeoutJob:
        def result(self, timeout=None):
            assert timeout == vendor_report.BIGQUERY_QUERY_TIMEOUT_SECONDS
            raise TimeoutError("deadline exceeded")

    class SuccessJob:
        def result(self, timeout=None):
            assert timeout == vendor_report.BIGQUERY_QUERY_TIMEOUT_SECONDS
            return EmptyResult()

    class RetryClient:
        def __init__(self):
            self.calls = 0

        def query(self, query, **kwargs):
            self.calls += 1
            assert kwargs["timeout"] == vendor_report.BIGQUERY_QUERY_TIMEOUT_SECONDS
            return TimeoutJob() if self.calls == 1 else SuccessJob()

    client = RetryClient()
    result = _query_to_dataframe(client, "SELECT organization", stage="bill lookup")

    assert client.calls == 2
    assert result.empty


def test_query_timeout_and_authentication_errors_are_actionable():
    class TimeoutJob:
        def result(self, timeout=None):
            raise TimeoutError("deadline exceeded")

    class TimeoutClient:
        def query(self, query, **kwargs):
            return TimeoutJob()

    with pytest.raises(RuntimeError, match="bill lookup.*after one retry") as timeout:
        _query_to_dataframe(
            TimeoutClient(),
            "SELECT 1",
            stage="bill lookup",
        )
    assert "timed out after one retry" in user_facing_report_error(
        timeout.value
    ).lower()

    class AuthJob:
        def result(self, timeout=None):
            raise RefreshError("Reauthentication is needed")

    class AuthClient:
        def query(self, query, **kwargs):
            return AuthJob()

    with pytest.raises(RuntimeError, match="Google authentication expired") as auth:
        _query_to_dataframe(
            AuthClient(),
            "SELECT 1",
            stage="payment lookup",
        )
    assert "gcloud auth application-default login" in user_facing_report_error(
        auth.value
    )


def test_safe_exception_details_redacts_credentials():
    details = safe_exception_details(
        RuntimeError("query failed; access_token=secret; Authorization:Bearer abc123")
    )

    assert "secret" not in details
    assert "abc123" not in details
    assert "[redacted]" in details


def test_generated_report_is_in_memory_and_has_expected_name(monkeypatch):
    datasets = _sample_datasets()
    metadata = {
        "queries": build_vendor_reconciliation_queries(
            "project-1",
            {"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 31)},
        ),
        "filters": {"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 31)},
    }
    monkeypatch.setattr(
        vendor_report,
        "fetch_vendor_report_data",
        lambda **kwargs: (datasets, metadata),
    )

    result = generate_vendor_transactions_report(
        date(2026, 1, 1),
        date(2026, 7, 31),
        report_mode=REPORT_MODE_ANALYSIS,
    )
    reopened = load_workbook(BytesIO(result["report_bytes"]), data_only=False)

    assert result["report_path"] is None
    assert result["report_name"].startswith(
        "Vendor_Reconciliation_Analysis_20260101_20260731_"
    )
    assert result["report_name"].endswith(".xlsx")
    assert reopened.sheetnames == WORKBOOK_SHEETS
    assert set(result["row_counts"]) == set(WORKBOOK_SHEETS)
    assert result["row_counts"]["Vendor Payments"] == 1


def test_generated_input_report_has_expected_name_and_sheets(monkeypatch):
    datasets = _sample_input_datasets()
    metadata = {
        "queries": build_reconciliation_input_queries(
            "project-1",
            {
                "start_date": date(2026, 1, 1),
                "end_date": date(2026, 7, 31),
                "bank_upload_id": "upload-full-period",
            },
        ),
        "filters": {"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 31)},
        "bank_upload": {
            "upload_id": "upload-full-period",
            "source_file": "statement.xlsx",
            "bank_source": "Uploaded Bank",
            "coverage_start_date": date(2026, 1, 1),
            "coverage_end_date": date(2026, 7, 31),
        },
    }
    monkeypatch.setattr(
        vendor_report,
        "fetch_reconciliation_input_data",
        lambda **kwargs: (datasets, metadata),
    )

    result = generate_vendor_transactions_report(
        date(2026, 1, 1),
        date(2026, 7, 31),
        report_mode=REPORT_MODE_INPUT,
    )
    reopened = load_workbook(BytesIO(result["report_bytes"]), data_only=False)

    assert result["report_name"].startswith(
        "Vendor_Reconciliation_Input_20260101_20260731_"
    )
    assert result["report_name"].endswith(".xlsx")
    assert reopened.sheetnames == INPUT_WORKBOOK_SHEETS
    assert set(result["row_counts"]) == set(INPUT_WORKBOOK_SHEETS)
    assert result["row_counts"]["Zoho Payment Batch"] == 3


def test_generated_report_creates_missing_output_folder(monkeypatch, tmp_path):
    datasets = _sample_input_datasets()
    metadata = {
        "queries": build_reconciliation_input_queries(
            "project-1",
            {
                "start_date": date(2026, 1, 1),
                "end_date": date(2026, 7, 31),
                "bank_upload_id": "upload-full-period",
            },
        ),
        "filters": {"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 31)},
        "bank_upload": {
            "upload_id": "upload-full-period",
            "source_file": "statement.xlsx",
            "bank_source": "Uploaded Bank",
            "coverage_start_date": date(2026, 1, 1),
            "coverage_end_date": date(2026, 7, 31),
        },
    }
    monkeypatch.setattr(
        vendor_report,
        "fetch_reconciliation_input_data",
        lambda **kwargs: (datasets, metadata),
    )
    destination = tmp_path / "new" / "reports"

    result = generate_vendor_transactions_report(
        date(2026, 1, 1),
        date(2026, 7, 31),
        report_mode=REPORT_MODE_INPUT,
        destination_folder=destination,
    )

    assert result["report_path"].exists()
    assert result["report_path"].parent == destination


def test_locked_output_error_is_safe_and_actionable(monkeypatch, tmp_path):
    datasets = _sample_input_datasets()
    metadata = {
        "queries": build_reconciliation_input_queries(
            "project-1",
            {
                "start_date": date(2026, 1, 1),
                "end_date": date(2026, 7, 31),
                "bank_upload_id": "upload-full-period",
            },
        ),
        "filters": {"start_date": date(2026, 1, 1), "end_date": date(2026, 7, 31)},
        "bank_upload": {
            "upload_id": "upload-full-period",
            "source_file": "statement.xlsx",
            "bank_source": "Uploaded Bank",
            "coverage_start_date": date(2026, 1, 1),
            "coverage_end_date": date(2026, 7, 31),
        },
    }
    monkeypatch.setattr(
        vendor_report,
        "fetch_reconciliation_input_data",
        lambda **kwargs: (datasets, metadata),
    )

    def locked_write(self, data):
        raise PermissionError("file is locked")

    monkeypatch.setattr(Path, "write_bytes", locked_write)

    with pytest.raises(RuntimeError, match="Close any open copy"):
        generate_vendor_transactions_report(
            date(2026, 1, 1),
            date(2026, 7, 31),
            report_mode=REPORT_MODE_INPUT,
            destination_folder=tmp_path,
        )


def test_invalid_input_date_range_is_rejected_before_querying():
    with pytest.raises(ValueError, match="Start date"):
        vendor_report.fetch_reconciliation_input_data(
            date(2026, 7, 31),
            date(2026, 7, 1),
            client=object(),
        )
